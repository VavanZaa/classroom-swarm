from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from lxml import etree

import build_indonesian_research_report as base


MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
MATHML_NS = "http://www.w3.org/1998/Math/MathML"
_MATH_TRANSFORM = etree.XSLT(etree.parse(str(MML2OMML)))


def body(doc: Document, text: str, lead: str | None = None) -> None:
    base.add_body(doc, text, bold_lead=lead)


def heading(doc: Document, text: str, level: int = 1) -> None:
    base.add_heading(doc, text, level)


def table(doc: Document, headers, rows, widths, font_size=9.0) -> None:
    base.add_table(doc, headers, rows, widths, font_size=font_size)


def equation(doc: Document, inner_mathml: str) -> None:
    """Add one centered native Word equation converted from presentation MathML."""
    mathml = etree.fromstring(
        f'<math xmlns="{MATHML_NS}" display="block"><mstyle displaystyle="true">'
        f'{inner_mathml}</mstyle></math>'.encode("utf-8")
    )
    transformed = _MATH_TRANSFORM(mathml)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph._p.append(etree.fromstring(etree.tostring(transformed.getroot())))


def build_humane_report(figures: dict[str, Path], means: dict[str, dict[str, float]]) -> None:
    doc = Document()
    base.setup_styles(doc)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(7)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.92)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.16)
    base.add_page_number(section.footer.paragraphs[0])

    # Cover: define the central idea before introducing technical names.
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(24)
    title.add_run("Laporan cara kerja simulasi kelas dengan kecerdasan artifisial")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Kecerdasan artifisial selanjutnya disebut AI. Model adalah kumpulan aturan yang dijalankan komputer. "
        "Simulasi adalah percobaan digital yang menjalankan model untuk membandingkan tiga keadaan kelas "
        "tanpa menempatkan siswa nyata pada risiko."
    )

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(10)
    run = meta.add_run(
        "Classroom Swarm V3 dan eksperimen E8, yaitu rangkaian 23.000 percobaan komputer\n"
        "Aturan dan hasil dibekukan pada September 2026\n"
        "Ditulis untuk guru dan perancang kebijakan sekolah"
    )
    base.set_run_font(run, size=11, color="4D4D4D")

    base.add_figure(
        doc,
        figures["anatomy"],
        "Gambar sampul. Setiap lingkaran adalah satu siswa tiruan. Setiap garis menunjukkan bahwa dua siswa dapat berinteraksi.",
        "Kelas tiruan berisi 30 siswa dalam tiga kelompok. Garis menghubungkan siswa yang dapat berinteraksi. Layanan AI berada di luar jaringan siswa.",
        width=6.5,
    )
    body(
        doc,
        "Model ini paling berguna untuk menanyakan apa yang mungkin terjadi jika aturan tertentu diterapkan. "
        "Model ini belum dapat meramalkan berapa banyak siswa nyata yang akan memakai AI atau seberapa besar perubahan nilai di sekolah tertentu, "
        "karena belum dibandingkan dengan data siswa nyata.",
        lead="Kesimpulan utama  ",
    )
    doc.add_page_break()

    heading(doc, "Ringkasan untuk pembaca", 1)
    body(
        doc,
        "Dalam laporan ini, model berarti kumpulan aturan yang dijalankan oleh komputer. Aturan itu menggambarkan kelas tiruan berisi 30 siswa. "
        "Setiap siswa memilih satu cara untuk menyelesaikan tugas, menerima hasil, lalu mengingat apakah pilihannya membantu. Komputer mengulang proses ini "
        "agar kita dapat melihat pola yang sering terjadi."
    )
    body(
        doc,
        "Tiga keadaan kelas dibandingkan. Keadaan pertama tidak menyediakan AI. Keadaan kedua menyediakan AI yang langsung memberikan jawaban lengkap. "
        "Keadaan ketiga menyediakan tutor AI yang memberi petunjuk bertahap dan tetap meminta siswa berusaha. Semua siswa boleh mengakses AI pada dua keadaan terakhir, "
        "tetapi setiap siswa tetap memilih apakah akan menggunakannya."
    )
    body(
        doc,
        "Di dalam model, AI yang memberi jawaban lengkap menaikkan nilai tugas pada saat bantuan tersedia, tetapi menurunkan hasil ketika siswa harus menjawab tanpa AI. "
        "Tutor AI lebih jarang dipilih, tetapi menghasilkan pengetahuan dan hasil ujian yang lebih baik. Arah temuan ini sejalan dengan penelitian yang relevan [1, 2]. "
        "Angkanya tetap merupakan hasil kelas tiruan, bukan perkiraan untuk sekolah nyata."
    )

    heading(doc, "Arti tingkat keyakinan dalam laporan", 2)
    body(
        doc,
        "Tingkat keyakinan menjelaskan seberapa jauh suatu hasil dapat dipercaya untuk tujuan tertentu. Nilai tinggi tidak berarti bahwa model sudah benar untuk semua sekolah."
    )
    table(
        doc,
        ["Bagian yang dinilai", "Keyakinan", "Cara membacanya"],
        [
            ["Perhitungan komputer", "Tinggi", "Aturan dikunci, berkas diperiksa, dan 222 pengujian perangkat lunak lulus."],
            ["Perbandingan tiga kebijakan di dalam model", "Cukup tinggi", "Ketiga keadaan memakai kelas dan kejadian acak yang sama. Perbedaannya hanya layanan AI."],
            ["Arah perbedaan antara jawaban langsung dan tutor", "Sedang", "Arah hasil bertahan pada beberapa aturan pilihan dan jumlah tugas yang diuji."],
            ["Persentase penggunaan dan besar perubahan hasil", "Rendah", "Angka sangat berubah ketika aturan pilihan siswa diubah dan belum disesuaikan dengan data nyata."],
            ["Ramalan untuk sekolah tertentu", "Rendah", "Model belum memuat data sekolah, percakapan AI nyata, atau perubahan strategi guru."],
        ],
        [1.75, 1.1, 3.95],
    )

    heading(doc, "Istilah yang perlu dipahami lebih dahulu", 1)
    body(doc, "Istilah berikut digunakan berulang kali. Penjelasannya diberikan di sini agar bagian berikutnya dapat dibaca tanpa latar belakang analisis data.")
    table(
        doc,
        ["Istilah", "Arti dalam laporan"],
        [
            ["Model", "Kumpulan aturan yang menggambarkan cara siswa memilih, belajar, dan berinteraksi."],
            ["Simulasi", "Percobaan digital yang menjalankan aturan model pada kelas tiruan."],
            ["Data sintetis", "Angka yang dihasilkan komputer dari simulasi. Angka ini bukan catatan siswa nyata."],
            ["Skenario", "Satu susunan aturan yang sengaja dipilih untuk menjawab pertanyaan tertentu."],
            ["Parameter", "Angka yang mengatur kuat atau lemahnya suatu proses, misalnya biaya usaha atau besarnya tambahan belajar."],
            ["Kalibrasi", "Proses menyesuaikan parameter agar mendekati data nyata yang relevan."],
            ["Validasi", "Pemeriksaan apakah model cukup sesuai untuk tujuan yang dinyatakan."],
            ["Agen", "Satu siswa tiruan yang mempunyai keadaan dan aturan pilihan sendiri."],
            ["Simpul", "Lingkaran yang mewakili seorang siswa pada gambar jaringan."],
            ["Hubungan atau garis", "Kesempatan bagi dua siswa untuk berdiskusi atau meminta jawaban."],
            ["Perilaku yang muncul", "Pola kelas yang terbentuk dari banyak keputusan siswa dan tidak ditetapkan sejak awal."],
            ["Replikasi", "Satu kali percobaan lengkap pada satu kelas tiruan."],
        ],
        [1.65, 5.1],
        font_size=9.0,
    )

    heading(doc, "Pertanyaan penelitian", 1)
    body(
        doc,
        "Penelitian ini menanyakan bagaimana layanan AI mengubah pilihan siswa, hasil belajar, dan bantuan antarsiswa ketika semua siswa memiliki akses, tetapi tidak semuanya memilih untuk memakai AI."
    )
    base.add_bullets(
        doc,
        [
            "Apakah siswa mulai memakai AI tanpa diberi label sebagai pengguna AI?",
            "Apakah jawaban lengkap dari AI membuat nilai tugas terlihat lebih tinggi daripada kemampuan saat AI tidak tersedia?",
            "Apakah tutor AI dapat membantu belajar walaupun siswa harus mengeluarkan usaha?",
            "Apakah penggunaan AI mengurangi diskusi dan bantuan antarsiswa?",
            "Aturan mana yang paling banyak mengubah hasil model?",
        ],
    )

    heading(doc, "Cara kelas digambarkan", 1)
    heading(doc, "Siswa dan keadaan awal", 2)
    body(
        doc,
        "Satu lingkaran pada gambar mewakili satu siswa tiruan. Setiap kelas berisi 30 siswa. Lingkaran tidak merujuk kepada siswa nyata dan tidak menyimpan data pribadi."
    )
    body(
        doc,
        "Setiap siswa memiliki beberapa catatan internal. Penjelasan sederhana diberikan lebih dahulu; lambang teknis dicantumkan sesudahnya agar pembaca dapat mencocokkannya dengan kode program."
    )
    table(
        doc,
        ["Catatan pada siswa", "Arti sederhana", "Nama dalam model"],
        [
            ["Pengetahuan", "Peluang dasar untuk menjawab benar tanpa bantuan", "K"],
            ["Catatan keberhasilan yang dirasakan", "Rata-rata bergerak dari nilai tugas sebelumnya; bukan ukuran psikologis", "C"],
            ["Efisiensi belajar", "Seberapa banyak pengetahuan dapat bertambah setelah belajar", "alpha"],
            ["Biaya usaha", "Seberapa berat suatu kegiatan dirasakan oleh siswa", "kappa"],
            ["Nilai pilihan", "Catatan pribadi tentang seberapa menguntungkan suatu pilihan", "Q"],
            ["Kepercayaan kepada sumber", "Perkiraan apakah seorang teman dapat dan mau memberikan jawaban", "Aij"],
        ],
        [2.25, 3.4, 1.15],
    )
    body(
        doc,
        "Pengetahuan awal, efisiensi belajar, dan biaya usaha diacak secara terpisah dari posisi siswa dalam jaringan. Karena itu, model tidak menciptakan kelompok tetap seperti siswa rajin, penyalin, pengguna AI, atau perantara jawaban."
    )

    heading(doc, "Pilihan yang tersedia bagi siswa", 2)
    body(doc, "Pada setiap tugas, siswa memilih satu dari empat tindakan berikut. Pilihan menggunakan AI hanya tersedia pada kelas yang menyediakan AI.")
    table(
        doc,
        ["Tindakan", "Apa yang dilakukan siswa", "Kode program"],
        [
            ["Mengerjakan sendiri", "Siswa berusaha menjawab dan belajar dari pengerjaan.", "S"],
            ["Berdiskusi", "Siswa memilih seorang teman dan belajar melalui diskusi yang disederhanakan.", "D"],
            ["Meminta jawaban teman", "Siswa meminta jawaban dari teman yang dianggap mampu dan bersedia membantu.", "C"],
            ["Menggunakan AI", "Siswa memakai layanan AI yang tersedia pada kebijakan tersebut.", "A"],
        ],
        [1.7, 4.2, 0.9],
    )

    heading(doc, "Cara siswa memilih", 2)
    body(
        doc,
        "Setelah siswa mencoba suatu tindakan, model mencatat hasil yang diterima. Catatan ini disebut nilai pilihan. Pada tugas berikutnya, siswa biasanya memilih tindakan yang memiliki nilai pribadi tertinggi. "
        "Aturan ini menggambarkan gagasan sederhana bahwa pengalaman yang menguntungkan cenderung diulangi. Model memakai bentuk komputasi dari pembelajaran penguatan yang dikenal sebagai Q-learning [11]. "
        "Penggunaan aturan ini tidak berarti siswa nyata menghitung rumus Q secara sadar."
    )
    body(
        doc,
        "Agar siswa tetap dapat mencoba pilihan lain, model memberi peluang kecil sebesar enam persen untuk memilih secara acak dari semua tindakan yang tersedia. Peluang mencoba ini disebut eksplorasi. "
        "Eksplorasi bersifat netral: AI, diskusi, mengerjakan sendiri, dan meminta jawaban mendapat kesempatan yang sama."
    )
    body(
        doc,
        "Pada tugas pertama saja, ada peluang sepuluh persen untuk meminta jawaban teman. Peluang awal ini disebut benih. Benih diperlukan karena tindakan tidak dapat muncul jika sejak awal peluangnya benar-benar nol. "
        "Benih dimatikan setelah tugas pertama. Model juga tidak menambah daya tarik suatu tindakan hanya karena tindakan itu populer di antara teman."
    )

    heading(doc, "Cara komputer menghitung peluang pilihan", 2)
    body(
        doc,
        "Bagian ini menunjukkan rumus yang benar-benar dipakai program. Penjelasan lambang diberikan sebelum rumus. Angka desimal dalam rumus memakai titik karena mengikuti penulisan kode komputer."
    )
    table(
        doc,
        ["Lambang", "Arti"],
        [
            ["i", "Siswa yang sedang membuat pilihan."],
            ["a", "Salah satu tindakan yang tersedia bagi siswa."],
            ["F_i", "Daftar tindakan yang dapat dipilih siswa i pada tugas tersebut."],
            ["n_i", "Jumlah tindakan dalam F_i. Nilainya empat jika AI dan tindakan sosial tersedia."],
            ["Q_i(a)", "Nilai yang diingat siswa i dari pengalaman sebelumnya dengan tindakan a."],
            ["U_i(a)", "Nilai yang dipakai komputer saat membandingkan tindakan."],
            ["P_i(a)", "Peluang siswa i memilih tindakan a."],
            ["epsilon", "Peluang eksplorasi, yaitu peluang mencoba secara acak. Nilai utama adalah 0,06."],
            ["B_i dan b_i", "Kumpulan tindakan dengan nilai U tertinggi dan jumlah tindakan di dalam kumpulan itu."],
            ["A_ij", "Kepercayaan pribadi siswa i bahwa teman j dapat dan mau memberikan jawaban."],
            ["w_ij", "Bobot kesempatan interaksi antara siswa i dan j. Semua bobot sama pada perbandingan utama E8."],
            ["kappa_i", "Biaya usaha yang dirasakan siswa i."],
        ],
        [1.55, 5.2],
        font_size=8.8,
    )

    heading(doc, "Peluang pada tugas pertama", 2)
    body(
        doc,
        "Pada tugas pertama, AI belum tersedia. Siswa yang mempunyai setidaknya satu hubungan sosial mendapat peluang 10 persen untuk meminta jawaban. Sisa peluang dibagi rata antara mengerjakan sendiri dan berdiskusi. Lambang s0 adalah nilai benih sebesar 0,10."
    )
    equation(
        doc,
        "<mrow><mi>P</mi><mo>(</mo><mi>C</mi><mo>)</mo><mo>=</mo><msub><mi>s</mi><mn>0</mn></msub><mo>=</mo><mn>0.10</mn><mo>,</mo>"
        "<mspace width='0.5em'/><mi>P</mi><mo>(</mo><mi>S</mi><mo>)</mo><mo>=</mo><mi>P</mi><mo>(</mo><mi>D</mi><mo>)</mo><mo>=</mo>"
        "<mfrac><mrow><mn>1</mn><mo>-</mo><msub><mi>s</mi><mn>0</mn></msub></mrow><mn>2</mn></mfrac><mo>=</mo><mn>0.45</mn><mo>,</mo>"
        "<mspace width='0.5em'/><mi>P</mi><mo>(</mo><mi>A</mi><mo>)</mo><mo>=</mo><mn>0</mn></mrow>",
    )
    body(
        doc,
        "Siswa tanpa hubungan sosial hanya dapat mengerjakan sendiri, sehingga peluang S sama dengan satu. Benih 10 persen hanya berlaku pada tugas pertama dan tidak menjadi sifat tetap siswa."
    )

    heading(doc, "Nilai tindakan mulai tugas kedua", 2)
    body(
        doc,
        "Sebelum memilih, siswa memperkirakan peluang memperoleh jawaban dari teman. N_i adalah kumpulan teman yang terhubung dengan siswa i. Peluang memilih teman j sebagai sasaran permintaan dinyatakan dengan pi_ij."
    )
    equation(
        doc,
        "<mrow><msub><mi>π</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo>"
        "<mfrac><mrow><msub><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><msub><mi>A</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub></mrow>"
        "<mrow><msub><mi>Σ</mi><mrow><mi>k</mi><mo>∈</mo><msub><mi>N</mi><mi>i</mi></msub></mrow></msub>"
        "<msub><mi>w</mi><mrow><mi>i</mi><mi>k</mi></mrow></msub><msub><mi>A</mi><mrow><mi>i</mi><mi>k</mi></mrow></msub></mrow></mfrac></mrow>",
    )
    body(doc, "Perkiraan akses keseluruhan, ditulis Â_i, adalah rata-rata kepercayaan kepada teman setelah peluang sasaran di atas diperhitungkan.")
    equation(
        doc,
        "<mrow><msub><mi>Â</mi><mi>i</mi></msub><mo>=</mo><msub><mi>Σ</mi>"
        "<mrow><mi>j</mi><mo>∈</mo><msub><mi>N</mi><mi>i</mi></msub></mrow></msub>"
        "<msub><mi>π</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><msub><mi>A</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub></mrow>",
    )
    body(
        doc,
        "Komputer kemudian menghitung nilai pembanding U. Mengerjakan sendiri, berdiskusi, dan menggunakan AI memakai nilai pengalaman Q secara langsung. Meminta jawaban mendapat penyesuaian kecil jika akses yang diperkirakan lebih baik atau lebih buruk daripada kepercayaan awal A0, yaitu 0,50."
    )
    equation(
        doc,
        "<mrow><msubsup><mi>U</mi><mi>i</mi><mi>S</mi></msubsup><mo>=</mo><msubsup><mi>Q</mi><mi>i</mi><mi>S</mi></msubsup><mo>,</mo>"
        "<mspace width='0.45em'/><msubsup><mi>U</mi><mi>i</mi><mi>D</mi></msubsup><mo>=</mo><msubsup><mi>Q</mi><mi>i</mi><mi>D</mi></msubsup><mo>,</mo>"
        "<mspace width='0.45em'/><msubsup><mi>U</mi><mi>i</mi><mi>C</mi></msubsup><mo>=</mo><msubsup><mi>Q</mi><mi>i</mi><mi>C</mi></msubsup>"
        "<mo>+</mo><msub><mi>κ</mi><mi>i</mi></msub><mo>(</mo><msub><mi>Â</mi><mi>i</mi></msub><mo>-</mo><msub><mi>A</mi><mn>0</mn></msub><mo>)</mo><mo>,</mo>"
        "<mspace width='0.45em'/><msubsup><mi>U</mi><mi>i</mi><mi>A</mi></msubsup><mo>=</mo><msubsup><mi>Q</mi><mi>i</mi><mi>A</mi></msubsup></mrow>",
    )
    body(
        doc,
        "B_i berisi tindakan dengan nilai U tertinggi. Pada aturan utama, 94 persen peluang diberikan kepada tindakan terbaik dan 6 persen dibagi rata untuk mencoba semua tindakan yang tersedia. Jika beberapa tindakan sama-sama terbaik, bagian 94 persen dibagi rata di antara tindakan tersebut."
    )
    equation(
        doc,
        "<mrow><msub><mi>P</mi><mi>i</mi></msub><mo>(</mo><mi>a</mi><mo>)</mo><mo>=</mo>"
        "<mfrac><mi>ε</mi><msub><mi>n</mi><mi>i</mi></msub></mfrac><mo>+</mo>"
        "<mfrac><mrow><mn>1</mn><mo>-</mo><mi>ε</mi></mrow><msub><mi>b</mi><mi>i</mi></msub></mfrac>"
        "<mtext> jika </mtext><mi>a</mi><mo>∈</mo><msub><mi>B</mi><mi>i</mi></msub></mrow>",
    )
    equation(
        doc,
        "<mrow><msub><mi>P</mi><mi>i</mi></msub><mo>(</mo><mi>a</mi><mo>)</mo><mo>=</mo>"
        "<mfrac><mi>ε</mi><msub><mi>n</mi><mi>i</mi></msub></mfrac>"
        "<mtext> jika </mtext><mi>a</mi><mo>∉</mo><msub><mi>B</mi><mi>i</mi></msub></mrow>",
    )
    body(
        doc,
        "Contoh berikut hanya menjelaskan perhitungan, bukan hasil rata-rata E8. Misalkan empat tindakan tersedia, epsilon bernilai 0,06, biaya usaha siswa 0,15, perkiraan akses 0,70, dan A0 bernilai 0,50. Nilai meminta jawaban naik dari 0,58 menjadi 0,61. Mengerjakan sendiri tetap memiliki nilai tertinggi, yaitu 0,66."
    )
    table(
        doc,
        ["Tindakan", "Nilai pengalaman Q", "Penyesuaian", "Nilai U", "Peluang dipilih"],
        [
            ["Mengerjakan sendiri", "0,66", "0", "0,66", "95,5%"],
            ["Berdiskusi", "0,62", "0", "0,62", "1,5%"],
            ["Meminta jawaban", "0,58", "+0,03", "0,61", "1,5%"],
            ["Menggunakan AI", "0,64", "0", "0,64", "1,5%"],
        ],
        [1.7, 1.35, 1.1, 1.0, 1.2],
        font_size=8.8,
    )
    body(
        doc,
        "Komputer mengambil satu angka acak antara nol dan satu, lalu memilih tindakan sesuai rentang peluang pada tabel. Karena itu, tindakan terbaik biasanya dipilih, tetapi tindakan lain tetap dapat dicoba."
    )

    heading(doc, "Aturan peluang alternatif untuk pemeriksaan ketahanan", 2)
    body(
        doc,
        "Perbandingan utama E8 memakai aturan nilai tertinggi di atas. Pemeriksaan tambahan memakai aturan lunak yang memberi peluang lebih besar kepada tindakan bernilai tinggi tanpa hampir selalu memilih satu tindakan. Aturan ini disebut softmax. Tau adalah suhu pilihan; nilai tau yang lebih besar membuat peluang tindakan lebih merata. U_max adalah nilai U tertinggi di antara tindakan yang tersedia."
    )
    equation(
        doc,
        "<mrow><msub><mi>P</mi><mi>i</mi></msub><mo>(</mo><mi>a</mi><mo>)</mo><mo>=</mo>"
        "<mfrac><mi>ε</mi><msub><mi>n</mi><mi>i</mi></msub></mfrac><mo>+</mo><mo>(</mo><mn>1</mn><mo>-</mo><mi>ε</mi><mo>)</mo>"
        "<mfrac><mrow><mi>exp</mi><mo>(</mo><mfrac><mrow><msubsup><mi>U</mi><mi>i</mi><mi>a</mi></msubsup><mo>-</mo><msub><mi>U</mi><mi>max</mi></msub></mrow><mi>τ</mi></mfrac><mo>)</mo></mrow>"
        "<mrow><msub><mi>Σ</mi><mrow><mi>b</mi><mo>∈</mo><msub><mi>F</mi><mi>i</mi></msub></mrow></msub>"
        "<mi>exp</mi><mo>(</mo><mfrac><mrow><msubsup><mi>U</mi><mi>i</mi><mi>b</mi></msubsup><mo>-</mo><msub><mi>U</mi><mi>max</mi></msub></mrow><mi>τ</mi></mfrac><mo>)</mo></mrow></mfrac></mrow>",
    )
    body(
        doc,
        "Hasil penggunaan AI berubah besar ketika aturan nilai tertinggi diganti dengan softmax. Karena itu, laporan menilai arah perbandingan kebijakan lebih dapat dipercaya daripada persentase penggunaan yang tepat."
    )

    heading(doc, "Cara siswa meminta dan membagikan jawaban", 2)
    body(
        doc,
        "Seorang siswa hanya dapat meminta jawaban dari teman yang terhubung dengannya. Ia lebih cenderung memilih teman yang sebelumnya terbukti dapat dan mau membantu. "
        "Kepercayaan itu berubah hanya setelah keduanya benar-benar berinteraksi; siswa tidak menerima daftar reputasi global."
    )
    body(
        doc,
        "Teman yang memiliki jawaban boleh menerima atau menolak permintaan. Kesediaannya dapat meningkat jika ia pernah menerima bantuan dari siswa yang kini meminta. "
        "Setiap siswa dapat membagikan paling banyak dua jawaban dalam satu tugas. Penelitian tentang hubungan timbal balik mendukung kemungkinan bahwa pengalaman menerima bantuan memengaruhi kemauan membantu [10]. "
        "Batas dua jawaban dan kekuatan hubungan timbal balik tetap merupakan pilihan skenario, bukan angka hasil pengukuran sekolah."
    )
    body(
        doc,
        "Peluang berbagi dihitung dari tiga angka. p0 adalah kesediaan dasar sebesar 0,80. Beta_R adalah kekuatan hubungan timbal balik sebesar 2. A_ji adalah kepercayaan siswa j kepada siswa i berdasarkan permintaan j sebelumnya. Lambang p_j menuju i berarti peluang siswa j berbagi kepada siswa i. Fungsi logistic mengubah hasil perhitungan menjadi peluang antara nol dan satu, sedangkan logit mengubah peluang dasar ke skala perhitungan."
    )
    equation(
        doc,
        "<mrow><msub><mi>p</mi><mrow><mi>j</mi><mo>→</mo><mi>i</mi></mrow></msub><mo>=</mo><mi>logistic</mi><mo>(</mo>"
        "<mi>logit</mi><mo>(</mo><msup><msub><mi>p</mi><mi>share</mi></msub><mn>0</mn></msup><mo>)</mo><mo>+</mo>"
        "<msub><mi>β</mi><mi>R</mi></msub><mo>(</mo><msub><mi>A</mi><mrow><mi>j</mi><mi>i</mi></mrow></msub><mo>-</mo><msub><mi>A</mi><mn>0</mn></msub><mo>)</mo><mo>)</mo></mrow>",
    )
    body(
        doc,
        "Jika A_ji masih sama dengan nilai awal 0,50, peluang berbagi tetap 0,80. Keberhasilan atau kegagalan permintaan langsung kemudian memperbarui kepercayaan peminta. Y_ij bernilai satu jika jawaban benar-benar diterima dan nol jika tidak. Laju pembaruan eta_A bernilai 0,20."
    )
    equation(
        doc,
        "<mrow><msup><msub><mi>A</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>′</mo></msup><mo>=</mo>"
        "<msub><mi>A</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>+</mo><msub><mi>η</mi><mi>A</mi></msub>"
        "<mo>(</mo><msub><mi>Y</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>-</mo><msub><mi>A</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>)</mo></mrow>",
    )
    body(
        doc,
        "Sebagai contoh, kepercayaan 0,50 naik menjadi 0,60 setelah satu permintaan berhasil dan turun menjadi 0,40 setelah satu permintaan gagal. Hanya siswa yang meminta yang memperbarui catatannya."
    )

    heading(doc, "Tiga keadaan kelas yang dibandingkan", 2)
    body(doc, "Layanan AI diperlakukan sebagai alat dari luar kelas. AI tidak dijadikan siswa atau teman di dalam jaringan.")
    table(
        doc,
        ["Keadaan kelas", "Layanan yang diterima", "Belajar dalam skenario utama", "Usaha"],
        [
            ["Tanpa AI", "Siswa hanya dapat bekerja sendiri, berdiskusi, atau meminta jawaban teman", "Mengikuti aturan belajar manusia dalam V3", "Bergantung pada tindakan"],
            ["AI jawaban langsung", "AI memberikan jawaban lengkap dan selalu benar dalam skenario utama", "Tambahan belajar sepersepuluh dari mengerjakan sendiri; sama dengan permintaan jawaban yang berhasil", "Tidak ada biaya usaha dalam skenario utama"],
            ["AI tutor terkendali", "AI memberi petunjuk bertahap dan tidak langsung memberikan jawaban akhir", "Tambahan belajar sama dengan mengerjakan sendiri dalam skenario utama", "Biaya usaha bernilai satu"],
        ],
        [1.35, 2.55, 2.15, 0.85],
        font_size=8.8,
    )
    body(
        doc,
        "Jika AI digambar sebagai simpul pusat yang terhubung kepada semua siswa, gambar akan menyamakan akses teknologi dengan pertemanan. AI juga otomatis tampak sebagai perantara terbesar. "
        "Karena itu, E8 menempatkan AI di luar jaringan siswa dan mencatat penggunaannya sebagai tindakan."
    )

    heading(doc, "Bagaimana tindakan mengubah pengetahuan dan hasil tugas", 2)
    body(
        doc,
        "Rumus berikut dipakai setelah tindakan dipilih. Semua perubahan memakai keadaan pada awal tugas. Fungsi clip membatasi pengetahuan agar tetap berada antara nol dan satu."
    )
    table(
        doc,
        ["Lambang", "Arti"],
        [
            ["K_i dan K_i'", "Pengetahuan siswa i sebelum dan sesudah tugas."],
            ["p_i", "Peluang siswa i menjawab satu butir dengan benar setelah tindakannya."],
            ["alpha_i", "Efisiensi belajar siswa i."],
            ["K_j", "Pengetahuan teman j yang dipilih untuk berdiskusi."],
            ["chi_C", "Bagian belajar yang tersisa setelah berhasil meminta jawaban; nilainya 0,10."],
            ["m_U dan m_T", "Pengali belajar untuk AI jawaban langsung dan tutor terkendali."],
            ["r_U", "Ketepatan jawaban AI langsung; nilainya satu pada skenario utama."],
            ["s_T", "Kekuatan bantuan segera dari tutor; nilainya 0,25 pada skenario utama."],
            ["G_i", "Nilai tugas siswa i, dihitung dari sepuluh butir soal."],
            ["c_i", "Biaya usaha tindakan: satu untuk usaha biasa, nol untuk bantuan tanpa usaha pada skenario utama."],
            ["R_i", "Hasil yang dirasakan setelah nilai tugas dikurangi biaya usaha."],
            ["eta_Q dan eta_C", "Laju pembaruan nilai pilihan dan catatan keberhasilan; keduanya bernilai 0,20."],
        ],
        [1.65, 5.1],
        font_size=8.8,
    )

    body(doc, "Jika siswa mengerjakan sendiri, pengetahuan bertambah sesuai ruang belajar yang masih tersisa. Peluang menjawab benar sama dengan pengetahuan setelah belajar.")
    equation(
        doc,
        "<mrow><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>=</mo><mi>clip</mi><mo>[</mo>"
        "<msub><mi>K</mi><mi>i</mi></msub><mo>+</mo><msub><mi>α</mi><mi>i</mi></msub><mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>K</mi><mi>i</mi></msub><mo>)</mo><mo>]</mo><mo>,</mo>"
        "<mspace width='0.5em'/><msubsup><mi>p</mi><mi>i</mi><mi>S</mi></msubsup><mo>=</mo><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup></mrow>",
    )
    body(
        doc,
        "Jika siswa berdiskusi dengan teman j, tambahan belajar lebih besar hanya ketika pengetahuan teman lebih tinggi. Peluang jawaban benar memakai nilai yang lebih tinggi antara pengetahuan baru siswa dan pengetahuan teman."
    )
    equation(
        doc,
        "<mrow><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>=</mo><mi>clip</mi><mo>[</mo>"
        "<msub><mi>K</mi><mi>i</mi></msub><mo>+</mo><msub><mi>α</mi><mi>i</mi></msub><mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>K</mi><mi>i</mi></msub><mo>)</mo>"
        "<mo>(</mo><mn>1</mn><mo>+</mo><mi>max</mi><mo>(</mo><msub><mi>K</mi><mi>j</mi></msub><mo>-</mo><msub><mi>K</mi><mi>i</mi></msub><mo>,</mo><mn>0</mn><mo>)</mo><mo>)</mo><mo>]</mo><mo>,</mo>"
        "<mspace width='0.5em'/><msubsup><mi>p</mi><mi>i</mi><mi>D</mi></msubsup><mo>=</mo><mi>max</mi><mo>(</mo><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>,</mo><msub><mi>K</mi><mi>j</mi></msub><mo>)</mo></mrow>",
    )
    body(
        doc,
        "Jika permintaan jawaban berhasil, siswa memakai jawaban yang diterima dan hanya memperoleh sepersepuluh tambahan belajar dari mengerjakan sendiri. Jika permintaan gagal, pengetahuan tidak bertambah dan siswa menjawab pada saat batas waktu dengan peluang benar sebesar K_i."
    )
    equation(
        doc,
        "<mrow><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>=</mo><mi>clip</mi><mo>[</mo>"
        "<msub><mi>K</mi><mi>i</mi></msub><mo>+</mo><msub><mi>χ</mi><mi>C</mi></msub><msub><mi>α</mi><mi>i</mi></msub>"
        "<mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>K</mi><mi>i</mi></msub><mo>)</mo><mo>]</mo>"
        "<mtext> jika permintaan berhasil</mtext></mrow>",
    )
    body(doc, "Jika siswa memakai AI jawaban langsung, m_U mengatur belajar yang tidak disengaja dan r_U mengatur peluang jawaban AI benar.")
    equation(
        doc,
        "<mrow><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>=</mo><mi>clip</mi><mo>[</mo>"
        "<msub><mi>K</mi><mi>i</mi></msub><mo>+</mo><msub><mi>m</mi><mi>U</mi></msub><msub><mi>α</mi><mi>i</mi></msub>"
        "<mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>K</mi><mi>i</mi></msub><mo>)</mo><mo>]</mo><mo>,</mo>"
        "<mspace width='0.5em'/><msubsup><mi>p</mi><mi>i</mi><mi>A</mi></msubsup><mo>=</mo><msub><mi>r</mi><mi>U</mi></msub></mrow>",
    )
    body(doc, "Jika siswa memakai tutor terkendali, m_T mengatur belajar yang bertahan dan s_T mengatur bantuan segera saat mengerjakan tugas.")
    equation(
        doc,
        "<mrow><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>=</mo><mi>clip</mi><mo>[</mo>"
        "<msub><mi>K</mi><mi>i</mi></msub><mo>+</mo><msub><mi>m</mi><mi>T</mi></msub><msub><mi>α</mi><mi>i</mi></msub>"
        "<mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>K</mi><mi>i</mi></msub><mo>)</mo><mo>]</mo><mo>,</mo>"
        "<mspace width='0.5em'/><msubsup><mi>p</mi><mi>i</mi><mi>T</mi></msubsup><mo>=</mo><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup>"
        "<mo>+</mo><msub><mi>s</mi><mi>T</mi></msub><mo>(</mo><mn>1</mn><mo>-</mo><msup><msub><mi>K</mi><mi>i</mi></msub><mo>′</mo></msup><mo>)</mo></mrow>",
    )
    body(
        doc,
        "Untuk setiap butir soal, komputer mengambil angka acak antara nol dan satu. Jawaban dinilai benar jika angka itu lebih kecil daripada peluang benar. x_iq bernilai satu untuk jawaban benar dan nol untuk jawaban salah. M_H adalah jumlah butir tugas, yaitu sepuluh."
    )
    equation(
        doc,
        "<mrow><msubsup><mi>G</mi><mi>i</mi><mi>H</mi></msubsup><mo>=</mo><mfrac><mn>1</mn><msub><mi>M</mi><mi>H</mi></msub></mfrac>"
        "<msubsup><mi>Σ</mi><mrow><mi>q</mi><mo>=</mo><mn>1</mn></mrow><msub><mi>M</mi><mi>H</mi></msub></msubsup>"
        "<msub><mi>x</mi><mrow><mi>i</mi><mi>q</mi></mrow></msub></mrow>",
    )
    body(
        doc,
        "Nilai tugas kemudian dikurangi biaya usaha untuk membentuk hasil yang dirasakan. Hanya nilai Q dari tindakan yang baru dipilih yang diperbarui. Jika hasil yang dirasakan lebih besar daripada Q lama, Q naik; jika lebih kecil, Q turun."
    )
    equation(
        doc,
        "<mrow><msub><mi>R</mi><mi>i</mi></msub><mo>=</mo><msubsup><mi>G</mi><mi>i</mi><mi>H</mi></msubsup>"
        "<mo>-</mo><msub><mi>κ</mi><mi>i</mi></msub><msub><mi>c</mi><mi>i</mi></msub><mo>,</mo><mspace width='0.6em'/>"
        "<msubsup><mi>Q</mi><mi>i</mi><mi>a</mi></msubsup><mo>′</mo><mo>=</mo><msubsup><mi>Q</mi><mi>i</mi><mi>a</mi></msubsup>"
        "<mo>+</mo><msub><mi>η</mi><mi>Q</mi></msub><mo>(</mo><msub><mi>R</mi><mi>i</mi></msub><mo>-</mo><msubsup><mi>Q</mi><mi>i</mi><mi>a</mi></msubsup><mo>)</mo></mrow>",
    )
    body(
        doc,
        "Catatan keberhasilan C_i mengikuti nilai tugas dengan laju eta_C. Catatan ini disimpan untuk pengukuran, tetapi tidak masuk ke rumus pilihan V3 atau E8."
    )
    equation(
        doc,
        "<mrow><msup><msub><mi>C</mi><mi>i</mi></msub><mo>′</mo></msup><mo>=</mo><msub><mi>C</mi><mi>i</mi></msub>"
        "<mo>+</mo><msub><mi>η</mi><mi>C</mi></msub><mo>(</mo><msubsup><mi>G</mi><mi>i</mi><mi>H</mi></msubsup><mo>-</mo><msub><mi>C</mi><mi>i</mi></msub><mo>)</mo></mrow>",
    )

    heading(doc, "Urutan kejadian dalam satu tugas", 2)
    base.add_numbered(
        doc,
        [
            "Model menyalin keadaan semua siswa pada awal tugas. Perhitungan dalam tugas itu memakai salinan yang sama agar siswa tidak memperoleh keuntungan hanya karena diproses lebih dahulu.",
            "Model menentukan tindakan yang tersedia. Mengerjakan sendiri selalu tersedia. Diskusi dan permintaan jawaban memerlukan setidaknya satu teman. AI tersedia bagi semua siswa dalam dua keadaan AI mulai tugas kedua.",
            "Pada tugas pertama, model memakai peluang benih 45 persen mengerjakan sendiri, 45 persen berdiskusi, dan 10 persen meminta jawaban. Pada tugas berikutnya, model menghitung U dan peluang pilihan dari pengalaman sebelumnya.",
            "Komputer mengambil angka acak untuk setiap siswa dan menetapkan satu tindakan sesuai peluang yang telah dihitung.",
            "Siswa yang mengerjakan sendiri menghasilkan jawabannya sendiri. Siswa yang berdiskusi memilih satu teman menurut bobot hubungan, lalu menghasilkan jawaban dengan peluang benar yang dipengaruhi hasil belajar dan pengetahuan teman.",
            "Siswa yang meminta jawaban memilih sasaran menurut bobot hubungan dan kepercayaan A_ij. Permintaan diproses paling banyak tiga putaran. Jawaban yang baru diterima hanya dapat diteruskan pada putaran berikutnya.",
            "Teman yang memiliki jawaban memutuskan apakah akan berbagi. Jika permintaan yang bersedia dilayani melebihi kapasitas, hanya dua permintaan dapat berhasil pada tugas itu.",
            "Siswa yang memilih AI menerima jawaban lengkap atau petunjuk tutor sesuai keadaan kelas. AI tidak meminta jawaban melalui jaringan siswa.",
            "Model menilai sepuluh butir tugas, menghitung perubahan pengetahuan, biaya usaha, hasil yang dirasakan, dan catatan keberhasilan.",
            "Setelah semua hasil selesai, siswa memperbarui Q untuk tindakan yang baru dipilih. Kepercayaan A_ij diperbarui hanya untuk permintaan langsung. Semua perubahan kemudian menjadi keadaan awal tugas berikutnya.",
            "Setelah setiap lima tugas, siswa mengerjakan ujian tanpa bantuan sebanyak 30 butir. Peluang benar sama dengan pengetahuan K. Ujian hanya mengukur hasil dan tidak mengubah keadaan siswa dalam E8.",
        ],
    )

    heading(doc, "Cara hubungan antarsiswa digambarkan", 1)
    body(
        doc,
        "Garis abu-abu antara dua siswa berarti keduanya memiliki kesempatan tetap untuk berinteraksi. Garis ini bukan bukti persahabatan, kedekatan emosional, atau hubungan sebab-akibat."
    )
    body(
        doc,
        "Komputer membentuk tiga kelompok yang lebih rapat di dalam kelompok daripada di antara kelompok. Cara membentuk jaringan berkelompok ini dikenal sebagai stochastic block model [12]. "
        "Bentuk tersebut masuk akal untuk kelas tiruan yang memiliki lingkaran pertemanan, tetapi tidak membuktikan bahwa semua kelas nyata mempunyai bentuk yang sama."
    )
    body(
        doc,
        "Sebuah hubungan disebut berbobot jika beberapa pasangan siswa diberi peluang berinteraksi lebih besar daripada pasangan lain. Model mampu memakai bobot, tetapi perbandingan utama E8 memberi semua garis bobot yang sama. "
        "Bobot berdasarkan kedekatan nyata belum digunakan karena data kelas belum tersedia."
    )
    base.add_figure(
        doc,
        figures["anatomy"],
        "Gambar 1. Satu kelas tiruan yang dipakai dalam eksperimen.",
        "Tiga kelompok siswa terhubung oleh garis kontak. Ukuran lingkaran menunjukkan pengetahuan awal. Layanan AI berada di luar jaringan siswa.",
    )

    heading(doc, "Tiga jenis informasi jaringan", 2)
    body(doc, "Model menyimpan tiga jenis informasi. Ketiganya perlu dipisahkan karena mempunyai fungsi yang berbeda.")
    table(
        doc,
        ["Jenis informasi", "Bentuk", "Fungsi"],
        [
            ["Kesempatan kontak", "Garis dua arah yang tetap", "Menentukan pasangan siswa yang dapat berinteraksi. Dalam kode disebut G."],
            ["Kepercayaan kepada sumber", "Nilai berarah yang dimiliki oleh siswa peminta", "Membantu siswa memilih siapa yang akan dimintai jawaban. Dalam kode disebut Aij."],
            ["Aliran jawaban yang benar-benar terjadi", "Panah dari pemberi kepada penerima", "Mencatat rute bantuan setelah simulasi. Dalam analisis disebut H."],
        ],
        [1.9, 2.15, 2.75],
    )

    heading(doc, "Aturan yang dipasang dan perilaku yang muncul", 1)
    body(
        doc,
        "Aturan yang dipasang adalah keputusan peneliti yang ditulis ke dalam program. Perilaku yang muncul adalah pola yang terbentuk setelah banyak siswa mengikuti aturan tersebut. "
        "Istilah ilmiahnya adalah emergence. Sebagai contoh, model tidak memberi label siswa tertentu sebagai sumber jawaban. Seorang siswa baru disebut sumber setelah banyak jawaban benar-benar mengalir darinya."
    )
    table(
        doc,
        ["Bagian model", "Status", "Penjelasan sederhana"],
        [
            ["Tidak ada hadiah karena suatu tindakan sedang populer", "Aturan yang dipasang", "Siswa tidak menyukai AI atau penyalinan hanya karena banyak teman melakukannya."],
            ["Benih meminta jawaban pada tugas pertama", "Aturan yang dipasang", "Memberi kesempatan awal agar tindakan dapat muncul, lalu dimatikan."],
            ["Peluang mencoba enam persen", "Aturan yang dipasang", "Memberi kesempatan yang sama untuk mencoba semua tindakan yang tersedia."],
            ["Perbedaan pengetahuan, efisiensi, dan biaya usaha", "Aturan yang dipasang", "Perbedaan siswa diacak dan tidak ditempelkan pada posisi jaringan."],
            ["Jawaban AI selalu benar dan mudah", "Aturan skenario utama", "Ini adalah keadaan batas untuk menguji risiko, bukan gambaran layanan nyata."],
            ["Tutor menahan jawaban dan membantu belajar", "Aturan skenario utama", "Kekuatan bantuan diuji pada beberapa nilai."],
            ["Siapa yang terus memakai AI", "Perilaku yang muncul", "Terbentuk dari percobaan dan pengalaman hasil."],
            ["Siapa yang menjadi sumber jawaban", "Perilaku yang muncul", "Tidak ditentukan sebelum simulasi."],
            ["Berkurangnya bantuan manusia", "Perilaku yang muncul", "Terjadi ketika siswa mengganti tindakan manusia dengan AI."],
            ["Perbedaan antara nilai tugas dan ujian", "Hasil yang muncul dengan syarat", "Terbentuk dari aturan bantuan dan belajar yang dipasang."],
        ],
        [2.55, 1.55, 2.7],
        font_size=8.7,
    )

    heading(doc, "Cara eksperimen dilakukan", 1)
    heading(doc, "Perbandingan yang adil", 2)
    body(
        doc,
        "Satu replikasi adalah satu percobaan lengkap pada satu kelas tiruan. Dalam setiap replikasi, tiga keadaan kelas memakai siswa, jaringan, soal, dan kejadian acak yang sama. "
        "Hanya layanan AI yang berbeda. Cara ini disebut desain berpasangan. Tujuannya adalah mengurangi kemungkinan bahwa hasil berbeda hanya karena satu kelas kebetulan lebih mudah daripada kelas lain."
    )
    body(
        doc,
        "Pemakaian rangkaian kejadian acak yang sama pada beberapa kebijakan dikenal sebagai common random numbers. Nama teknis ini hanya menjelaskan alat perhitungan; gagasan sederhananya adalah membandingkan tiga versi dari kelas yang sama."
    )

    heading(doc, "Jumlah percobaan", 2)
    body(doc, "Satu kali menjalankan satu keadaan kelas disebut satu pelaksanaan simulasi. E8 terdiri atas 23.000 pelaksanaan yang dibagi menurut tujuan berikut.")
    table(
        doc,
        ["Kelompok percobaan", "Jumlah pelaksanaan", "Tujuan"],
        [
            ["Perbandingan utama", "1.200", "Empat ratus kelas dikalikan tiga keadaan kelas."],
            ["Pemeriksaan cara kerja", "7.400", "Mengubah ketepatan jawaban AI, kekuatan tutor, bantuan bertahap, dan biaya usaha."],
            ["Pemeriksaan ketahanan hasil", "14.400", "Mengubah aturan pilihan, pengaruh teman, bentuk jaringan, sifat siswa, dan jumlah tugas."],
            ["Total", "23.000", "Seluruh percobaan yang dilaporkan dalam E8."],
        ],
        [1.8, 1.15, 3.8],
    )

    heading(doc, "Rencana analisis dan pemeriksaan", 2)
    body(
        doc,
        "Pertanyaan, dugaan hasil, batas yang dianggap penting, dan cara mengelompokkan kebijakan ditulis sebelum hasil dilihat. Langkah ini disebut praregistrasi. "
        "Mesin V3 tidak diubah selama E8. Pemeriksaan menemukan tidak ada berkas yang hilang atau mempunyai sidik digital yang salah dari 32.615 berkas. Semua 222 pengujian perangkat lunak juga lulus. "
        "Pemeriksaan ini menunjukkan bahwa perhitungan dapat ditelusuri, tetapi tidak membuktikan bahwa perilaku siswa nyata sama dengan perilaku model."
    )

    heading(doc, "Pengulangan acak dan rentang ketidakpastian", 2)
    body(
        doc,
        "Simulasi diulang dengan variasi acak untuk melihat seberapa banyak hasil berubah dari satu kelas tiruan ke kelas lain. Cara pengulangan ini disebut metode Monte Carlo. "
        "Rata-rata 400 kelas lebih stabil daripada hasil satu kelas. Rentang ketidakpastian dalam E8 menunjukkan variasi di antara kelas tiruan tersebut. "
        "Rentang ini tidak menggambarkan seluruh siswa Indonesia karena tidak ada sampel siswa Indonesia yang masuk ke model."
    )

    heading(doc, "Hasil utama", 1)
    body(doc, "Semua angka pada bagian ini berasal dari simulasi. Nilai 100 berarti hasil maksimum pada skala model. Persentase penggunaan menunjukkan bagian tindakan pada lima tugas terakhir.")
    base.add_figure(
        doc,
        figures["primary"],
        "Gambar 2. Pilihan siswa dan hasil belajar rata-rata pada 400 kelas untuk setiap keadaan.",
        "Grafik kiri membandingkan pilihan siswa. Grafik kanan membandingkan nilai tugas, pengetahuan, ujian tanpa AI, dan selisih nilai tugas dengan ujian.",
    )
    labels = [
        ("pA_late", "Menggunakan AI"),
        ("pS_late", "Mengerjakan sendiri"),
        ("pD_late", "Berdiskusi"),
        ("pC_late", "Meminta jawaban teman"),
        ("grade_late", "Nilai tugas"),
        ("K_final", "Pengetahuan pada akhir simulasi"),
        ("exam_final", "Ujian tanpa AI"),
        ("grade_exam_gap", "Selisih nilai tugas dan ujian"),
    ]
    result_rows = [[label] + [f"{means[c][key] * 100:.1f}" for c in base.REGIME_ORDER] for key, label in labels]
    table(doc, ["Ukuran dalam poin", "Tanpa AI", "Jawaban langsung", "Tutor"], result_rows, [2.75, 1.15, 1.55, 1.15])
    body(
        doc,
        "Dibandingkan dengan keadaan tanpa AI, jawaban langsung menaikkan nilai tugas sebesar 4,8 poin. Pada saat yang sama, pengetahuan turun 3,3 poin dan hasil ujian tanpa AI turun 3,4 poin. "
        "Jarak antara nilai tugas dan ujian melebar 8,2 poin. Dengan kata lain, kinerja saat bantuan tersedia membaik, tetapi kemampuan tanpa bantuan memburuk di dalam model."
    )
    body(
        doc,
        "Tutor AI menaikkan nilai tugas sebesar 1,0 poin, pengetahuan 2,9 poin, dan ujian tanpa AI 2,9 poin dibandingkan dengan keadaan tanpa AI. Jarak antara nilai tugas dan ujian mengecil 1,9 poin. "
        "Jika tutor dibandingkan langsung dengan AI jawaban lengkap, keunggulan tutor pada pengetahuan dan ujian sekitar 6,2 poin."
    )
    body(
        doc,
        "Arah perbedaan antara tutor dan jawaban langsung bertahan ketika model memakai beberapa cara memilih tindakan dan ketika jumlah tugas dinaikkan dari 20 menjadi 40. "
        "Namun, persentase siswa yang memakai AI berubah besar ketika cara memilih tindakan diubah. Karena itu, angka penggunaan tidak layak dipakai sebagai ramalan."
    )
    base.add_figure(
        doc,
        figures["trajectory"],
        "Gambar 3. Perubahan penggunaan AI, pengetahuan, dan nilai tugas selama 20 tugas.",
        "Tiga grafik garis menunjukkan rata-rata penggunaan AI, pengetahuan, dan nilai tugas dari tugas pertama sampai tugas kedua puluh.",
    )

    heading(doc, "Cara membaca gambar aliran bantuan", 1)
    body(
        doc,
        "Gambar berikut memakai replikasi nomor 282. Replikasi ini dipilih karena hasilnya paling dekat dengan rata-rata 400 kelas pada sepuluh ukuran utama. "
        "Pemilihan dilakukan sebelum gambar dibuat agar peneliti tidak memilih kelas yang paling dramatis. Gambar tetap hanya menunjukkan satu kelas; kesimpulan kebijakan harus dibaca dari rata-rata 400 kelas."
    )
    base.add_figure(
        doc,
        figures["flows"],
        "Gambar 4. Aliran jawaban antarsiswa pada lima tugas terakhir.",
        "Ketiga panel memakai siswa dan hubungan kontak yang sama. Panah biru menunjukkan jawaban yang berhasil dipindahkan. Lingkaran besar menunjukkan siswa yang lebih sering menjadi sumber. Warna ungu menunjukkan penggunaan AI.",
    )
    base.add_bullets(
        doc,
        [
            "Garis abu-abu sama pada ketiga panel karena layanan AI tidak mengubah kesempatan kontak dasar.",
            "Panah biru berkurang ketika AI tersedia karena sebagian siswa beralih dari bantuan manusia ke layanan AI.",
            "Lingkaran besar menunjukkan peran sumber yang terbentuk dari kejadian nyata dalam simulasi, bukan label yang diberikan sejak awal.",
            "Warna ungu menunjukkan seberapa sering seorang siswa memakai AI. AI tidak digambar sebagai teman atau siswa pusat.",
        ],
    )
    body(
        doc,
        "Konsentrasi sumber menunjukkan apakah sebagian kecil siswa memasok sebagian besar jawaban. Tanpa AI, 20 persen sumber teratas memasok 62,5 persen jawaban. "
        "Bagian itu naik menjadi 81,8 persen pada AI jawaban langsung dan 75,6 persen pada tutor. Model juga menghitung indeks Gini, yaitu angka antara nol dan satu untuk mengukur ketimpangan pasokan. "
        "Nilainya naik dari 0,647 menjadi 0,772 dan 0,736. Hasil ini berarti bantuan manusia yang tersisa menjadi lebih bergantung pada sedikit sumber."
    )

    heading(doc, "Cara pengaruh dapat menyebar", 1)
    body(
        doc,
        "Pembelajaran sosial berarti seorang siswa belajar dari tindakan dan hasil yang ia lihat pada teman. Fitur ini tidak digunakan pada model utama. Dalam model utama, siswa mencoba AI melalui peluang eksplorasi, lalu mempertahankan atau meninggalkannya berdasarkan pengalaman sendiri. "
        "Karena itu, model utama tidak menunjukkan penyebaran penggunaan AI dari teman ke teman."
    )
    body(
        doc,
        "Sebagai pemeriksaan tambahan, E8 menjalankan versi yang membolehkan siswa mengamati tindakan dan hasil teman. Semua tindakan diperlakukan sama; AI tidak mendapat hadiah khusus. "
        "Pada versi ini, penggunaan AI jawaban langsung naik sekitar 25,9 poin dan hasil ujian turun sekitar 5,5 poin. Penggunaan tutor naik sekitar 12,9 poin dan ujian naik tipis 0,4 poin. "
        "Angka tersebut menunjukkan akibat dari aturan yang diuji, bukan laju penyebaran AI di kelas nyata. Penelitian menunjukkan bahwa belajar dari orang lain dapat berguna, tetapi keberhasilannya bergantung pada informasi dan lingkungan [9]."
    )
    body(
        doc,
        "Dua teman dapat memakai AI karena saling memengaruhi, tetapi juga karena minat, guru, atau akses mereka sama. Kecenderungan orang yang mirip untuk saling terhubung disebut homofili. "
        "Data pengamatan biasa sulit memisahkan homofili, pilihan teman, dan pengaruh langsung tanpa rancangan penelitian yang kuat [15]."
    )

    heading(doc, "Aturan yang paling banyak mengubah hasil", 1)
    body(doc, "Bagian ini membahas sensitivitas, yaitu seberapa besar hasil berubah ketika satu parameter atau aturan diganti.")
    table(
        doc,
        ["Aturan yang diubah", "Apa yang terjadi", "Makna bagi penelitian"],
        [
            ["Cara memilih tindakan", "Penggunaan AI langsung sekitar 31 persen ketika siswa selalu memilih nilai tertinggi. Angkanya menjadi 81 persen atau 56 persen ketika pilihan dibuat dengan peluang yang berbeda.", "Ini adalah pengaruh terbesar pada angka penggunaan. Data perilaku nyata diperlukan sebelum angka penggunaan dapat dipercaya."],
            ["Kekuatan belajar dari tutor", "Ketika kekuatan tutor dinaikkan dari 0,50 menjadi 1,50, pengetahuan naik sekitar 3,17 poin.", "Asumsi tentang kualitas belajar tutor sangat menentukan manfaat tutor."],
            ["Biaya usaha tutor", "Ketika biaya dinaikkan dari satu menjadi dua, penggunaan tutor turun sekitar 10,52 poin.", "Tutor yang terlalu membebani siswa mungkin efektif bagi pengguna, tetapi jarang dipilih."],
            ["Pengamatan terhadap teman", "Penggunaan AI jawaban langsung naik sekitar 25,9 poin dan kerugian ujian membesar.", "Pengaruh teman penting jika mekanisme ini benar-benar terjadi."],
            ["Jaringan berkelompok atau acak", "Perubahan rata-rata pada penggunaan AI dan hasil utama hampir nol.", "Semua siswa sudah memiliki akses AI, tetapi bentuk jaringan tetap penting bagi bantuan manusia."],
            ["Dua puluh atau empat puluh tugas", "Kerugian dari jawaban langsung dan keuntungan tutor tetap mempunyai arah yang sama.", "Temuan utama tidak hanya muncul sesaat pada akhir tugas kedua puluh."],
        ],
        [1.7, 2.9, 2.25],
        font_size=8.55,
    )

    heading(doc, "Dukungan penelitian untuk aturan yang dipasang", 1)
    body(
        doc,
        "Penelitian sejawat dapat mendukung alasan memilih suatu mekanisme. Penelitian tersebut tidak otomatis membenarkan angka parameter model. Kolom terakhir menjelaskan batas dukungan untuk setiap klaim."
    )
    evidence_rows = [
        ["AI tanpa pembatas dapat menaikkan kinerja saat bantuan tersedia tetapi merugikan ujian tanpa bantuan", "Eksperimen acak pada hampir 1.000 siswa matematika sekolah menengah [1]", "Mendukung arah risiko. Tidak membenarkan ketepatan AI 100 persen, biaya nol, atau besar perubahan dalam E8."],
        ["Tutor yang memberi petunjuk, umpan balik, dan tetap meminta usaha dapat membantu belajar", "Penelitian tentang AI tutor, tutor manusia, bantuan bertahap, dan umpan balik [1, 2, 3, 4, 8]", "Mendukung mekanisme umum. Tidak membuktikan bahwa setiap tutor atau setiap nilai kekuatan tutor akan efektif."],
        ["Diskusi dan kegiatan belajar aktif dapat membantu belajar", "Meta-analisis 225 penelitian bidang sains, teknologi, teknik, dan matematika [5]", "Mendukung adanya jalur belajar dari kegiatan aktif. Tidak menentukan besar tambahan belajar dalam model."],
        ["Ujian tanpa bantuan memberi informasi berbeda dari nilai yang diperoleh dengan bantuan", "Eksperimen AI pendidikan dan penelitian tentang latihan mengingat kembali [1, 7]", "Mendukung pemisahan dua ukuran. Jumlah soal dan peluang benar dalam E8 tetap buatan."],
        ["Pengalaman hasil dapat mengubah nilai suatu pilihan", "Metode komputasi Q-learning [11]", "Mendukung cara komputer memperbarui nilai pilihan. Tidak membuktikan bahwa pikiran siswa bekerja dengan rumus yang sama."],
        ["Pengamatan terhadap teman dapat mengubah strategi", "Turnamen strategi pembelajaran sosial [9]", "Mendukung kemungkinan umum. Tidak menentukan kuatnya pengaruh teman atau laju penggunaan AI."],
        ["Hubungan timbal balik dapat membantu kerja sama", "Tinjauan ilmiah tentang resiprositas [10]", "Mendukung arah kemauan berbagi. Kapasitas dua jawaban dan kuatnya pengaruh belum disesuaikan dengan data kelas."],
        ["Kelas dapat digambarkan sebagai beberapa kelompok hubungan", "Penelitian tentang stochastic block model dan homofili [12, 16]", "Mendukung cara membuat jaringan tiruan. Jumlah kelompok, jumlah hubungan, dan jaringan yang tetap bukan fakta untuk sekolah tertentu."],
        ["Interaksi kelas dapat dipelajari sebagai jaringan", "Penelitian jaringan interaksi pada pembelajaran fisika [17]", "Mendukung penggunaan analisis jaringan. Hubungan yang ditemukan pada studi tersebut tidak membuktikan hubungan sebab-akibat dalam E8."],
        ["Kepercayaan kepada sumber berubah hanya setelah pengalaman langsung", "Belum ada validasi langsung untuk persamaan Aij", "Ini adalah aturan transparan untuk menghindari reputasi global buatan. Aturan perlu diuji dengan data permintaan bantuan."],
        ["Perbedaan siswa tidak dikaitkan dengan posisi jaringan", "Pilihan desain untuk mencegah tipe siswa ditetapkan sejak awal", "Membantu menjaga kemunculan peran. Bentuk sebaran sifat siswa masih belum didasarkan pada data nyata."],
        ["Catatan C mengikuti nilai tugas", "Tidak diperlakukan sebagai klaim psikologis", "C hanya catatan internal. Ia tidak boleh disebut kepercayaan diri nyata tanpa alat ukur yang telah divalidasi."],
    ]
    table(doc, ["Aturan atau klaim", "Dukungan penelitian", "Yang belum dibuktikan"], evidence_rows, [2.3, 2.25, 2.35], font_size=8.0)

    heading(doc, "Keterbatasan", 1)
    heading(doc, "Belum ada penyesuaian dengan data siswa nyata", 2)
    body(
        doc,
        "Ini adalah keterbatasan terbesar. Model belum mengetahui seberapa sering siswa Indonesia memakai AI, seberapa banyak mereka belajar dari satu interaksi, seberapa berat usaha dirasakan, atau bagaimana bantuan antarsiswa terbentuk. "
        "Karena itu, angka penggunaan 31 persen dan kenaikan pengetahuan 2,9 poin adalah hasil skenario, bukan perkiraan lapangan."
    )
    heading(doc, "Hubungan siswa tidak berubah dan belum diberi bobot nyata", 2)
    body(
        doc,
        "Kesempatan kontak tetap sama selama seluruh tugas. Model dapat mencatat siapa membantu siapa, tetapi bantuan itu tidak memperkuat atau melemahkan hubungan pada tugas berikutnya. "
        "Model juga belum memisahkan hubungan untuk diskusi dari hubungan untuk meminta jawaban."
    )
    heading(doc, "Pilihan dan proses belajar masih sederhana", 2)
    body(
        doc,
        "Satu tindakan mewakili proses yang dalam kenyataan dapat berlangsung lama. Model tidak menyimpan isi diskusi, tidak memuat deteksi guru terhadap penyalinan, dan tidak membuat tutor menilai kesalahpahaman melalui percakapan. "
        "Belajar dihitung segera setelah tindakan. Model belum memuat lupa, perpindahan pengetahuan ke topik lain, atau retensi jangka panjang."
    )
    heading(doc, "Dua layanan AI utama dibuat terlalu ideal", 2)
    body(
        doc,
        "AI jawaban langsung selalu benar dan tidak meminta usaha dalam skenario utama. Tutor selalu memiliki solusi yang telah diperiksa dan memberi umpan balik yang tepat. "
        "Layanan nyata dapat salah, berubah kualitas, memakai model yang berbeda, atau digunakan di luar sistem sekolah."
    )
    heading(doc, "Guru dan sekolah belum ikut menyesuaikan aturan", 2)
    body(
        doc,
        "Guru tidak mengubah tugas, penilaian, kelompok, atau aturan setelah melihat perilaku siswa. Model juga belum memuat privasi, disiplin, pemeriksaan plagiarisme, biaya perangkat, kesenjangan akses di rumah, atau pelatihan literasi AI."
    )
    heading(doc, "Bukti penelitian berasal dari konteks yang terbatas", 2)
    body(
        doc,
        "Penelitian utama yang digunakan sebagai pembanding berasal dari matematika sekolah menengah pada satu sekolah dan fisika perguruan tinggi. Tinjauan tutor cerdas untuk pendidikan dasar dan menengah menemukan hasil yang umumnya positif, tetapi berbeda-beda antarpenelitian [4]. "
        "Temuan tidak boleh dipindahkan langsung ke bahasa, seni, pendidikan vokasi, atau kebutuhan khusus."
    )
    heading(doc, "Pengaruh teman sulit dibuktikan dari pengamatan biasa", 2)
    body(
        doc,
        "Jika dua teman sama-sama memakai AI, penyebabnya belum tentu pengaruh langsung. Mereka mungkin memiliki minat, guru, atau akses yang sama. Model dapat memisahkan penyebab karena peneliti mengatur setiap aturan. Data sekolah nyata memerlukan rancangan yang lebih kuat untuk melakukan hal yang sama [15]."
    )

    heading(doc, "Apa yang dapat dan tidak dapat diwakili oleh data ini", 1)
    table(
        doc,
        ["Dapat digunakan untuk", "Belum dapat digunakan untuk"],
        [
            ["Memeriksa akibat logis dari aturan yang tertulis", "Meramalkan persentase pengguna AI di sekolah tertentu"],
            ["Membandingkan arah perubahan antara jawaban langsung, usaha, dan belajar", "Menentukan besar pengaruh sebab-akibat di dunia nyata"],
            ["Membandingkan tiga kebijakan pada kelas tiruan yang sama", "Menilai atau mendiagnosis seorang siswa"],
            ["Menentukan pola apa yang perlu dicari dalam data nyata", "Meramalkan nilai, kepercayaan diri, atau risiko menyontek seseorang"],
            ["Menguji apakah kesimpulan bertahan ketika aturan diubah", "Membuktikan bahwa pengaruh teman benar-benar terjadi di kelas"],
            ["Menentukan data minimum yang perlu dikumpulkan", "Membenarkan larangan atau kewajiban penggunaan AI secara otomatis"],
        ],
        [3.35, 3.35],
    )
    body(
        doc,
        "Cara paling aman untuk menyatakan hasil adalah sebagai kalimat bersyarat. Jika siswa belajar dari pengalaman seperti dalam model, jika AI jawaban langsung menggantikan usaha, dan jika tutor menghasilkan belajar yang setara dengan mengerjakan sendiri, maka arah hasil dalam laporan akan muncul. "
        "Tanpa kata jika, pernyataan tersebut akan terdengar seperti klaim tentang sekolah nyata, padahal data pendukungnya belum ada."
    )

    heading(doc, "Kemungkinan penggunaan", 1)
    table(
        doc,
        ["Tujuan", "Cara menggunakan hasil", "Batas yang harus dijaga"],
        [
            ["Membandingkan rancangan kebijakan", "Uji beberapa kekuatan tutor, biaya usaha, dan batas akses sebelum pilot sekolah.", "Gunakan rentang hasil, bukan satu angka ramalan."],
            ["Membantu merancang pembelajaran", "Pertahankan diskusi manusia, tugas tanpa AI, dan petunjuk yang tetap meminta usaha.", "Model tidak menentukan satu metode terbaik untuk semua mata pelajaran."],
            ["Menyiapkan pilot", "Pilih ukuran penggunaan AI, nilai dengan bantuan, ujian tanpa bantuan, usaha, dan aliran bantuan.", "Pilot memerlukan persetujuan etika dan perlindungan privasi."],
            ["Memantau keadaan kelas", "Amati apakah jarak nilai tugas dan ujian melebar atau diskusi manusia berkurang pada tingkat kelompok.", "Jangan menandai atau menghukum siswa berdasarkan model."],
            ["Menguji sistem AI yang lebih mandiri", "Gunakan V3 sebagai pembanding yang dapat diperiksa ketika menambahkan AI percakapan.", "Interaksi yang lebih realistis tidak menggantikan data lapangan."],
        ],
        [1.45, 3.25, 2.25],
        font_size=8.6,
    )

    heading(doc, "Langkah sebelum menambahkan sistem AI yang lebih mandiri", 1)
    body(
        doc,
        "Sistem AI yang lebih mandiri berarti model bahasa membuat respons baru selama simulasi, bukan hanya mengikuti nilai jawaban dan belajar yang sudah ditentukan. Tambahan ini dapat membuat percakapan lebih realistis, tetapi juga lebih sulit diperiksa."
    )
    base.add_numbered(
        doc,
        [
            "Pertahankan V3 dan E8 tanpa perubahan sebagai pembanding dasar.",
            "Biarkan aturan pilihan siswa tetap dapat diperiksa. Mulailah dengan mengganti kualitas respons tutor, bukan memberikan seluruh keputusan siswa kepada model bahasa.",
            "Periksa apakah komponen baru menghasilkan rentang hasil yang sama dengan E8 pada keadaan sederhana. Jika tidak, catat aturan baru yang menyebabkan perbedaan.",
            "Simpan catatan setiap masukan, respons, kesalahan, petunjuk, percobaan siswa, dan penggunaan jawaban akhir.",
            "Ulangi pengujian dengan beberapa model bahasa dan beberapa versi. Hasil yang hanya muncul pada satu model atau satu susunan perintah belum cukup stabil untuk rekomendasi sekolah.",
            "Pisahkan pemeriksaan teknis dari klaim pendidikan. Percakapan yang tampak nyata tidak otomatis membuat gambaran belajar siswa menjadi benar.",
        ],
    )

    heading(doc, "Data minimum untuk pemeriksaan pada masa depan", 1)
    body(
        doc,
        "Sekolah belum perlu mengumpulkan semua percakapan siswa. Jika pilot kelak dilakukan, data minimum berikut dapat dikumpulkan dalam bentuk kelompok atau dengan identitas samaran."
    )
    table(
        doc,
        ["Data", "Kegunaan", "Contoh pengukuran yang lebih aman"],
        [
            ["Penggunaan AI pada setiap tugas", "Mengetahui kapan siswa mencoba, melanjutkan, atau berhenti memakai AI", "Catatan jenis layanan tanpa menyimpan isi percakapan pada tahap awal"],
            ["Nilai dengan bantuan dan kuis tanpa bantuan", "Mengukur jarak antara kinerja saat dibantu dan kemampuan mandiri", "Kuis singkat berkala dengan aturan yang jelas"],
            ["Jenis bantuan AI", "Membedakan jawaban langsung, petunjuk, umpan balik, dan penjelasan", "Kategori layanan dan jumlah langkah bantuan"],
            ["Interaksi antarsiswa", "Menyesuaikan bentuk hubungan dan aliran bantuan", "Survei singkat atau pengamatan kelompok tentang siapa berdiskusi dengan siapa"],
            ["Waktu dan usaha", "Menyesuaikan biaya yang memengaruhi pilihan", "Waktu pengerjaan dan skala usaha singkat"],
            ["Pengetahuan awal dan daya ingat", "Menyesuaikan perbedaan awal dan perubahan belajar", "Tes awal dan tes ulang setelah jeda"],
            ["Persepsi diri yang diukur dengan alat sah", "Mengganti catatan C dengan ukuran psikologis yang benar", "Skala efikasi diri yang sesuai usia dan konteks"],
        ],
        [1.7, 2.65, 2.45],
        font_size=8.55,
    )
    body(
        doc,
        "Pengumpulan data memerlukan persetujuan sekolah dan, sesuai usia, persetujuan orang tua atau siswa. Sekolah perlu mengumpulkan sesedikit mungkin data, membatasi siapa yang dapat mengaksesnya, dan melaporkan hasil pada tingkat kelompok. "
        "Data hubungan antarsiswa sangat sensitif karena masih dapat mengungkap pergaulan meskipun nama telah dihapus."
    )

    heading(doc, "Penilaian akhir terhadap kekuatan model", 1)
    body(
        doc,
        "Model cukup kuat sebagai percobaan tentang cara kerja suatu mekanisme. Aturannya dapat diperiksa, siswa tidak diberi peran sosial sejak awal, popularitas AI tidak diberi hadiah, hubungan kontak dipisahkan dari aliran jawaban, dan tiga kebijakan dibandingkan pada kelas tiruan yang sama. "
        "Hasil juga telah diuji ketika beberapa aturan utama diubah."
    )
    body(
        doc,
        "Model belum cukup kuat sebagai alat ramalan atau alat penilaian siswa. Parameter belajar belum disesuaikan dengan data nyata, hubungan antarsiswa belum diukur, layanan tutor masih sederhana, dan persentase penggunaan sangat peka terhadap aturan pilihan. "
        "Klaim yang dapat dipertanggungjawabkan adalah arah perubahan dengan syarat tertentu, bukan angka penggunaan atau tingkat keberhasilan di sekolah."
    )
    body(
        doc,
        "Langkah penelitian berikutnya sebaiknya menguji satu tambahan pada satu waktu. Setiap tambahan perlu menghasilkan ukuran atau dugaan baru yang dapat diperiksa dengan data, bukan hanya membuat simulasi tampak lebih rumit."
    )

    heading(doc, "Referensi sejawat", 1)
    references = [
        ("1", "Bastani H, Bastani O, Sungu A, Ge H, Kabakci O, dan Mariman R. Generative AI without guardrails can harm learning. Proceedings of the National Academy of Sciences. 2025.", "https://doi.org/10.1073/pnas.2422633122"),
        ("2", "Kestin G, Miller K, Klales A, Milbourne T, dan kolega. AI tutoring outperforms in class active learning. Scientific Reports. 2025.", "https://doi.org/10.1038/s41598-025-97652-6"),
        ("3", "VanLehn K. The relative effectiveness of human tutoring intelligent tutoring systems and other tutoring systems. Educational Psychologist. 2011.", "https://doi.org/10.1080/00461520.2011.611369"),
        ("4", "Letourneau A dan kolega. A systematic review of AI driven intelligent tutoring systems in K to 12 education. npj Science of Learning. 2025.", "https://doi.org/10.1038/s41539-025-00320-7"),
        ("5", "Freeman S, Eddy SL, McDonough M, Smith MK, Okoroafor N, Jordt H, dan Wenderoth MP. Active learning increases student performance in science engineering and mathematics. PNAS. 2014.", "https://doi.org/10.1073/pnas.1319030111"),
        ("6", "Sweller J dan Cooper GA. The use of worked examples as a substitute for problem solving in learning algebra. Cognition and Instruction. 1985.", "https://doi.org/10.1207/s1532690xci0201_3"),
        ("7", "Roediger HL dan Karpicke JD. Test enhanced learning. Psychological Science. 2006.", "https://doi.org/10.1111/j.1467-9280.2006.01693.x"),
        ("8", "Hattie J dan Timperley H. The power of feedback. Review of Educational Research. 2007.", "https://doi.org/10.3102/003465430298487"),
        ("9", "Rendell L dan kolega. Why copy others. Insights from the social learning strategies tournament. Science. 2010.", "https://doi.org/10.1126/science.1184719"),
        ("10", "Nowak MA dan Sigmund K. Evolution of indirect reciprocity. Nature. 2005.", "https://doi.org/10.1038/nature04131"),
        ("11", "Watkins CJCH dan Dayan P. Q-learning. Machine Learning. 1992.", "https://doi.org/10.1007/BF00992698"),
        ("12", "Holland PW, Laskey KB, dan Leinhardt S. Stochastic blockmodels first steps. Social Networks. 1983.", "https://doi.org/10.1016/0378-8733(83)90021-7"),
        ("13", "Grimm V dan kolega. The ODD protocol for describing agent based and other simulation models. Journal of Artificial Societies and Social Simulation. 2020.", "https://doi.org/10.18564/jasss.4259"),
        ("14", "Epstein JM. Why model. Journal of Artificial Societies and Social Simulation. 2008.", "https://www.jasss.org/11/4/12.html"),
        ("15", "Shalizi CR dan Thomas AC. Homophily and contagion are generically confounded in observational social network studies. Sociological Methods and Research. 2011.", "https://doi.org/10.1177/0049124111404820"),
        ("16", "McPherson M, Smith Lovin L, dan Cook JM. Birds of a feather. Homophily in social networks. Annual Review of Sociology. 2001.", "https://doi.org/10.1146/annurev.soc.27.1.415"),
        ("17", "Bruun J dan Brewe E. Talking and learning physics. Physical Review Special Topics Physics Education Research. 2013.", "https://doi.org/10.1103/PhysRevSTPER.9.020109"),
    ]
    for num, citation, url in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.23)
        p.paragraph_format.first_line_indent = Inches(-0.23)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(f"[{num}] {citation} ")
        base.set_run_font(run, size=9.4)
        base.add_hyperlink(p, "DOI atau artikel", url)

    heading(doc, "Catatan agar penelitian dapat diperiksa ulang", 1)
    body(
        doc,
        "Laporan ini memakai hasil E8 yang sudah dibekukan. Angka utama berasal dari berkas ringkasan dan perjalanan hasil utama. Gambar jaringan memakai catatan mentah replikasi 282, yaitu kelas yang paling dekat dengan rata-rata gabungan. "
        "Laporan ini tidak menjalankan simulasi baru."
    )
    table(
        doc,
        ["Bagian", "Keterangan"],
        [
            ["Versi model", "Classroom Swarm V3 dengan eksperimen E8"],
            ["Jumlah pelaksanaan simulasi", "23.000"],
            ["Perbandingan utama", "400 replikasi untuk setiap keadaan kelas"],
            ["Kelas untuk gambar jaringan", "Replikasi 282 dipilih dengan aturan jarak terhadap rata-rata"],
            ["Penyesuaian dengan data sekolah", "Belum dilakukan"],
            ["Cara AI digambarkan", "Layanan dari luar jaringan siswa"],
        ],
        [2.25, 4.45],
    )

    doc.core_properties.title = "Laporan cara kerja simulasi kelas dengan kecerdasan artifisial"
    doc.core_properties.subject = "Penjelasan sederhana metodologi dan hasil Classroom Swarm V3 E8"
    doc.core_properties.author = "Classroom Swarm Research Project"
    doc.core_properties.keywords = "simulasi kelas jaringan siswa AI pendidikan"
    base.REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(base.OUT)


def main() -> None:
    base.REPORT_DIR.mkdir(parents=True, exist_ok=True)
    base.FIG_DIR.mkdir(parents=True, exist_ok=True)
    base.set_mpl()
    data = base.load_representative()
    means = base.primary_means()
    figures = {
        "anatomy": base.save_network_anatomy(data),
        "flows": base.save_flow_comparison(data),
        "primary": base.save_primary_results(means),
        "trajectory": base.save_trajectories(),
    }
    build_humane_report(figures, means)
    print(base.OUT)


if __name__ == "__main__":
    main()
