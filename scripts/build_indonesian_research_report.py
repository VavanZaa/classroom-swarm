"""Build the Indonesian E8 methodology report and its embedded figures.

The script reads only frozen E8 outputs. It does not rerun or mutate the model.
"""

from __future__ import annotations

import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable, Sequence

import matplotlib as mpl
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pyarrow.parquet as pq
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "docs" / "reports"
FIG_DIR = REPORT_DIR / "figures"
OUT = REPORT_DIR / "Laporan_Metodologi_Simulasi_Jaringan_Kelas_AI_E8.docx"
E8 = ROOT / "runs" / "E8"
REP = 282

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF4FA"
GRAY = "D9E2F3"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "7F8C8D"
TEXT = "202124"
RED = "C00000"
GREEN = "548235"
ORANGE = "C55A11"

REGIME_LABEL = {
    "E8-A-none": "Tanpa AI",
    "E8-A-unrestricted": "AI jawaban langsung",
    "E8-A-controlled": "AI tutor terkendali",
}
REGIME_ORDER = ["E8-A-none", "E8-A-unrestricted", "E8-A-controlled"]
REGIME_COLOR = {
    "E8-A-none": "#7A7A7A",
    "E8-A-unrestricted": "#C44E52",
    "E8-A-controlled": "#4C78A8",
}


def read_rows(path: Path) -> list[dict]:
    return pq.read_table(path).to_pylist()


def mean(values: Iterable[float | None]) -> float:
    xs = [float(x) for x in values if x is not None]
    return float(np.mean(xs)) if xs else float("nan")


def set_mpl() -> None:
    mpl.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 10,
            "axes.titlesize": 12,
            "axes.labelsize": 10,
            "axes.titleweight": "bold",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )


def load_representative() -> dict:
    rep_dir = E8 / "primary" / f"rep{REP:04d}"
    return {
        "nodes": read_rows(rep_dir / "graph_nodes.parquet"),
        "edges": read_rows(rep_dir / "graph_edges.parquet"),
        "traits": read_rows(rep_dir / "agent_traits.parquet"),
        "decisions": read_rows(rep_dir / "decisions.parquet"),
        "routes": read_rows(rep_dir / "routing_events.parquet"),
        "ai": read_rows(rep_dir / "ai_decisions.parquet"),
    }


def graph_for_run(data: dict, run_id: str) -> nx.Graph:
    g = nx.Graph()
    for r in data["nodes"]:
        if r["run_id"] == run_id:
            g.add_node(int(r["student_id"]), community=int(r["community_id"]), degree=int(r["degree"]))
    for r in data["edges"]:
        if r["run_id"] == run_id:
            g.add_edge(int(r["u"]), int(r["v"]), weight=float(r["weight"]))
    return g


def run_id_for(data: dict, condition: str) -> str:
    for r in data["nodes"]:
        if r["run_id"].startswith(condition + "/"):
            return str(r["run_id"])
    raise KeyError(condition)


def save_network_anatomy(data: dict) -> Path:
    path = FIG_DIR / "01_anatomi_jaringan_kelas.jpg"
    run_id = run_id_for(data, "E8-A-none")
    g = graph_for_run(data, run_id)
    traits = {
        int(r["student_id"]): float(r["K0"])
        for r in data["traits"]
        if r["run_id"] == run_id
    }
    pos = nx.spring_layout(g, seed=20260921, k=0.52, iterations=250, weight="weight")
    community_colors = {0: "#4C78A8", 1: "#F28E2B", 2: "#59A14F"}

    fig, ax = plt.subplots(figsize=(10.8, 7.1))
    nx.draw_networkx_edges(g, pos, ax=ax, edge_color="#B7BEC8", width=1.3, alpha=0.75)
    for community in sorted({g.nodes[n]["community"] for n in g.nodes}):
        nodes = [n for n in g.nodes if g.nodes[n]["community"] == community]
        sizes = [260 + 780 * traits[n] for n in nodes]
        nx.draw_networkx_nodes(
            g,
            pos,
            nodelist=nodes,
            node_size=sizes,
            node_color=community_colors[community],
            edgecolors="white",
            linewidths=1.7,
            label=f"Kelompok {community + 1}",
            ax=ax,
        )
    labels = {n: f"S{n + 1:02d}" for n in g.nodes}
    nx.draw_networkx_labels(g, pos, labels=labels, font_size=7.6, font_color="white", font_weight="bold", ax=ax)

    # Show that a private belief and a transfer have direction even though the social edge is undirected.
    u, v = sorted(g.edges())[9]
    x1, y1 = pos[u]
    x2, y2 = pos[v]
    ax.annotate(
        "",
        xy=(x2, y2),
        xytext=(x1, y1),
        arrowprops=dict(arrowstyle="-|>", color="#7B2CBF", lw=3.0, shrinkA=18, shrinkB=18),
        zorder=5,
    )
    midx, midy = (x1 + x2) / 2, (y1 + y2) / 2
    ax.text(
        midx - 0.16,
        midy - 0.14,
        "keyakinan pribadi atau aliran jawaban\nmemiliki arah",
        color="#5A189A",
        fontsize=8.5,
        ha="center",
        va="bottom",
        bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="#B8A1D9", alpha=0.95),
    )

    ax.text(
        1.03,
        0.83,
        "Layanan AI",
        transform=ax.transAxes,
        ha="center",
        va="center",
        fontsize=11,
        fontweight="bold",
        color="#17365D",
        bbox=dict(boxstyle="round,pad=0.55", fc="#EAF2F8", ec="#2F75B5", lw=1.5),
        clip_on=False,
    )
    ax.annotate(
        "akses tersedia bagi semua\nnamun AI bukan simpul siswa",
        xy=(1.03, 0.75),
        xytext=(1.03, 0.58),
        xycoords=ax.transAxes,
        textcoords=ax.transAxes,
        ha="center",
        va="center",
        fontsize=8.5,
        color="#17365D",
        arrowprops=dict(arrowstyle="->", color="#2F75B5", lw=1.2),
        clip_on=False,
    )
    ax.set_title("Anatomi jaringan kelas yang dipakai model", loc="left", pad=14)
    ax.text(
        0.0,
        1.01,
        "30 lingkaran adalah siswa; garis abu-abu adalah peluang kontak sosial yang tetap",
        transform=ax.transAxes,
        fontsize=9.5,
        color="#4D4D4D",
        va="bottom",
    )
    ax.legend(loc="lower left", frameon=False, ncol=3, fontsize=9, bbox_to_anchor=(0.0, -0.02))
    ax.text(
        0.0,
        -0.075,
        "Ukuran lingkaran menunjukkan pengetahuan awal K0 pada replikasi 282. Warna hanya menunjukkan kelompok jaringan.",
        transform=ax.transAxes,
        fontsize=8.5,
        color="#5A5A5A",
    )
    ax.set_axis_off()
    fig.subplots_adjust(left=0.02, right=0.83, top=0.90, bottom=0.10)
    fig.savefig(path, dpi=260, bbox_inches="tight", pil_kwargs={"quality": 93})
    plt.close(fig)
    return path


def successful_flows(data: dict, run_id: str, hmin: int = 16) -> Counter:
    counts: Counter = Counter()
    for r in data["routes"]:
        if r["run_id"] == run_id and int(r["homework_index"]) >= hmin and bool(r["transfer_success"]):
            # The answer moves from the target/holder to the requester.
            counts[(int(r["target_id"]), int(r["requester_id"]))] += 1
    return counts


def ai_use_rates(data: dict, run_id: str, hmin: int = 16) -> dict[int, float]:
    used: defaultdict[int, list[int]] = defaultdict(list)
    for r in data["ai"]:
        if r["run_id"] == run_id and int(r["homework_index"]) >= hmin:
            used[int(r["student_id"])].append(int(bool(r["used_A"])))
    return {i: float(np.mean(v)) for i, v in used.items()}


def save_flow_comparison(data: dict) -> Path:
    path = FIG_DIR / "02_perbandingan_aliran_bantuan.jpg"
    base_run = run_id_for(data, "E8-A-none")
    g = graph_for_run(data, base_run)
    pos = nx.spring_layout(g, seed=20260921, k=0.52, iterations=250, weight="weight")

    fig, axes = plt.subplots(1, 3, figsize=(13.5, 5.3))
    cmap = mpl.colors.LinearSegmentedColormap.from_list("ai", ["#F7F7F7", "#8E63CE", "#3F007D"])
    norm = mpl.colors.Normalize(vmin=0, vmax=1)
    panel_stats = []

    for ax, condition in zip(axes, REGIME_ORDER):
        run_id = run_id_for(data, condition)
        flows = successful_flows(data, run_id)
        rates = ai_use_rates(data, run_id)
        supplied = Counter()
        for (source, _receiver), count in flows.items():
            supplied[source] += count

        nx.draw_networkx_edges(g, pos, ax=ax, edge_color="#D8DDE4", width=0.65, alpha=0.62)
        if flows:
            flow_g = nx.DiGraph()
            for (source, receiver), count in flows.items():
                flow_g.add_edge(source, receiver, count=count)
            widths = [0.7 + 0.75 * math.sqrt(flow_g[u][v]["count"]) for u, v in flow_g.edges]
            nx.draw_networkx_edges(
                flow_g,
                pos,
                ax=ax,
                edge_color="#1F77B4",
                width=widths,
                alpha=0.72,
                arrows=True,
                arrowsize=9,
                arrowstyle="-|>",
                connectionstyle="arc3,rad=0.08",
                min_source_margin=9,
                min_target_margin=9,
            )
        node_colors = [cmap(norm(rates.get(n, 0.0))) for n in g.nodes]
        node_sizes = [150 + 85 * supplied[n] for n in g.nodes]
        nx.draw_networkx_nodes(
            g,
            pos,
            ax=ax,
            node_size=node_sizes,
            node_color=node_colors,
            edgecolors="#34495E",
            linewidths=0.65,
        )
        nx.draw_networkx_labels(g, pos, labels={n: str(n + 1) for n in g.nodes}, font_size=6.5, ax=ax)
        human_transfers = sum(flows.values())
        ai_actions = sum(
            1
            for r in data["ai"]
            if r["run_id"] == run_id and int(r["homework_index"]) >= 16 and bool(r["used_A"])
        )
        panel_stats.append((condition, human_transfers, ai_actions))
        ax.set_title(REGIME_LABEL[condition], pad=8)
        ax.text(
            0.5,
            -0.04,
            f"transfer manusia berhasil {human_transfers}   |   tindakan AI {ai_actions}",
            transform=ax.transAxes,
            ha="center",
            fontsize=8.4,
            color="#4D4D4D",
        )
        ax.set_axis_off()

    sm = mpl.cm.ScalarMappable(cmap=cmap, norm=norm)
    cax = fig.add_axes([0.22, 0.075, 0.58, 0.035])
    cbar = fig.colorbar(sm, cax=cax, orientation="horizontal")
    cbar.set_label("Proporsi penggunaan AI per siswa pada lima tugas terakhir")
    fig.suptitle("Aliran bantuan pada kelas yang sama di tiga kebijakan", fontsize=15, fontweight="bold", x=0.02, ha="left")
    fig.text(
        0.02,
        0.91,
        "Garis biru berarah adalah transfer jawaban manusia yang berhasil. Simpul lebih besar berarti siswa lebih sering menjadi sumber.",
        fontsize=9.5,
        color="#4D4D4D",
        ha="left",
    )
    fig.subplots_adjust(left=0.02, right=0.98, top=0.78, bottom=0.20, wspace=0.08)
    fig.savefig(path, dpi=260, bbox_inches="tight", pil_kwargs={"quality": 93})
    plt.close(fig)
    return path


def primary_means() -> dict[str, dict[str, float]]:
    rows = read_rows(E8 / "derived" / "primary_summaries.parquet")
    outcomes = [
        "pA_late",
        "pS_late",
        "pD_late",
        "pC_late",
        "grade_late",
        "K_final",
        "exam_final",
        "grade_exam_gap",
        "C_final",
        "source_gini",
        "source_top20",
    ]
    out: dict[str, dict[str, float]] = {}
    for condition in REGIME_ORDER:
        subset = [r for r in rows if r["experiment_id"] == condition]
        out[condition] = {name: mean(r[name] for r in subset) for name in outcomes}
    return out


def save_primary_results(means: dict[str, dict[str, float]]) -> Path:
    path = FIG_DIR / "03_hasil_utama_e8.jpg"
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.2), gridspec_kw={"width_ratios": [1.15, 1]})
    x = np.arange(4)
    width = 0.23
    action_keys = ["pS_late", "pD_late", "pC_late", "pA_late"]
    action_labels = ["Mengerjakan\nsendiri", "Diskusi", "Menyalin dari\nteman", "Menggunakan\nAI"]
    for idx, condition in enumerate(REGIME_ORDER):
        vals = [means[condition][k] * 100 for k in action_keys]
        axes[0].bar(x + (idx - 1) * width, vals, width, label=REGIME_LABEL[condition], color=REGIME_COLOR[condition])
    axes[0].set_xticks(x, action_labels)
    axes[0].set_ylabel("Persentase tindakan pada lima tugas terakhir")
    axes[0].set_ylim(0, 46)
    axes[0].grid(axis="y", color="#E5E5E5", lw=0.8)
    axes[0].legend(frameon=False, fontsize=8.5, loc="upper right")
    axes[0].set_title("Perubahan pilihan siswa", loc="left")

    outcome_keys = ["grade_late", "K_final", "exam_final", "grade_exam_gap"]
    outcome_labels = ["Nilai tugas", "Pengetahuan akhir", "Ujian tanpa AI", "Selisih nilai tugas\ndan ujian"]
    y = np.arange(len(outcome_keys))
    offsets = [-0.18, 0, 0.18]
    for offset, condition in zip(offsets, REGIME_ORDER):
        vals = [means[condition][k] * 100 for k in outcome_keys]
        axes[1].scatter(vals, y + offset, s=62, color=REGIME_COLOR[condition], label=REGIME_LABEL[condition], zorder=3)
        for xx, yy in zip(vals, y + offset):
            axes[1].text(xx + 1.0, yy, f"{xx:.1f}", va="center", fontsize=8, color=REGIME_COLOR[condition])
    axes[1].set_yticks(y, outcome_labels)
    axes[1].invert_yaxis()
    axes[1].set_xlim(0, 100)
    axes[1].set_xlabel("Poin dari skala 0 sampai 100")
    axes[1].grid(axis="x", color="#E5E5E5", lw=0.8)
    axes[1].set_title("Hasil belajar dan kinerja", loc="left")
    axes[1].legend(frameon=False, fontsize=8.3, loc="lower right")
    fig.suptitle("Hasil utama dari 400 kelas berpasangan per kebijakan", fontsize=15, fontweight="bold", x=0.02, ha="left")
    fig.text(
        0.02,
        0.915,
        "Angka adalah rata-rata simulasi E8, bukan estimasi untuk sekolah nyata.",
        fontsize=9.5,
        color="#4D4D4D",
        ha="left",
    )
    fig.subplots_adjust(left=0.08, right=0.97, top=0.82, bottom=0.20, wspace=0.33)
    fig.savefig(path, dpi=260, bbox_inches="tight", pil_kwargs={"quality": 93})
    plt.close(fig)
    return path


def save_trajectories() -> Path:
    path = FIG_DIR / "04_perjalanan_perilaku_dan_belajar.jpg"
    rows = read_rows(E8 / "derived" / "primary_trajectories.parquet")
    grouped: dict[tuple[str, int], list[dict]] = defaultdict(list)
    for r in rows:
        grouped[(r["experiment_id"], int(r["homework_index"]))].append(r)
    fig, axes = plt.subplots(1, 3, figsize=(12.7, 4.4), sharex=True)
    metrics = [("pA", "Penggunaan AI", "% tindakan"), ("K", "Pengetahuan", "skala 0 sampai 1"), ("grade", "Nilai tugas", "skala 0 sampai 1")]
    for ax, (metric, title, ylabel) in zip(axes, metrics):
        for condition in REGIME_ORDER:
            xs = list(range(1, 21))
            ys = [mean(r[metric] for r in grouped[(condition, h)]) for h in xs]
            if metric == "pA":
                ys = [100 * y for y in ys]
            ax.plot(xs, ys, color=REGIME_COLOR[condition], lw=2.2, label=REGIME_LABEL[condition])
        ax.axvspan(16, 20, color="#F2F2F2", zorder=0)
        ax.set_title(title, loc="left")
        ax.set_xlabel("Nomor tugas")
        ax.set_ylabel(ylabel)
        ax.set_xticks([1, 5, 10, 15, 20])
        ax.grid(color="#E6E6E6", lw=0.7)
        ax.set_ylim(bottom=0)
    axes[0].legend(frameon=False, fontsize=8.2, loc="upper left")
    fig.suptitle("Perjalanan rata-rata selama 20 tugas", fontsize=15, fontweight="bold", x=0.02, ha="left")
    fig.text(0.02, 0.91, "Area abu-abu adalah jendela akhir yang dipakai untuk ukuran perilaku akhir.", fontsize=9.3, color="#4D4D4D")
    fig.subplots_adjust(left=0.07, right=0.98, top=0.80, bottom=0.17, wspace=0.34)
    fig.savefig(path, dpi=260, bbox_inches="tight", pil_kwargs={"quality": 93})
    plt.close(fig)
    return path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=105, bottom=90, end=105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_col_widths(table, widths: Sequence[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool = False, color: str = TEXT, size: float = 9.2, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Arial"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float], font_size: float = 9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    repeat_header(hdr)
    prevent_row_split(hdr)
    for i, (cell, header) in enumerate(zip(hdr.cells, headers)):
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
    for ridx, row in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        for cidx, (cell, value) in enumerate(zip(cells, row)):
            align = WD_ALIGN_PARAGRAPH.CENTER if cidx > 0 and len(str(value)) < 24 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, str(value), size=font_size, align=align)
            if ridx % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
    set_col_widths(table, widths)
    set_table_borders(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def set_run_font(run, size=None, bold=None, color=None, italic=None, name="Arial") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE):
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(rfonts)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_run_font(run, size=8.5, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_picture_alt(run, description: str) -> None:
    for doc_pr in run._r.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)
        doc_pr.set("title", description[:120])


def add_figure(doc: Document, image_path: Path, caption: str, alt: str, width: float = 6.65) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    set_picture_alt(run, alt)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(9)
    cap.paragraph_format.keep_together = True
    r = cap.add_run(caption)
    set_run_font(r, size=9, italic=True, color="4D4D4D")


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.13
    p.paragraph_format.widow_control = True
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True, size=10.6)
    r = p.add_run(text)
    set_run_font(r, size=10.6)


def add_bullets(doc: Document, items: Sequence[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        p.add_run(item)
        for run in p.runs:
            set_run_font(run, size=10.4)


def add_numbered(doc: Document, items: Sequence[str]) -> None:
    # Use explicit numbering so each explanatory sequence restarts at 1.
    # Word/LibreOffice otherwise continue the previous list across sections.
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.add_run(f"{number}.\t{item}")
        for run in p.runs:
            set_run_font(run, size=10.4)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.6)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13
    for style_name, size in (("Title", 27), ("Subtitle", 13), ("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Subtitle"
        style.font.underline = False
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
    styles["Title"].paragraph_format.keep_with_next = True
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.color.rgb = RGBColor.from_string("4D4D4D")


def build_doc(figures: dict[str, Path], means: dict[str, dict[str, float]]) -> None:
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    # Keep body text clear of the page-number footer in both Word and LibreOffice.
    section.bottom_margin = Inches(0.92)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.16)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(24)
    p.add_run("Laporan metodologi simulasi jaringan kelas dengan akses AI")
    sp = doc.add_paragraph(style="Subtitle")
    sp.add_run("Model berbasis agen untuk membandingkan tanpa AI, AI jawaban langsung, dan AI tutor terkendali")
    mp = doc.add_paragraph()
    mp.paragraph_format.space_before = Pt(10)
    r = mp.add_run("Classroom Swarm V3 dan eksperimen E8\nStatus model beku September 2026\nLaporan untuk pendidik dan perancang kebijakan sekolah")
    set_run_font(r, size=11, color="4D4D4D")
    add_figure(
        doc,
        figures["anatomy"],
        "Gambar sampul. Satu contoh jaringan kelas yang benar-benar berasal dari keluaran E8.",
        "Jaringan kelas berisi 30 siswa dalam tiga kelompok. Garis menghubungkan siswa yang dapat berinteraksi. AI digambarkan terpisah karena bukan siswa.",
        width=6.5,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Kesimpulan utama")
    set_run_font(r, bold=True, size=11.2)
    p.add_run(
        "  Model ini berguna untuk menguji mekanisme dan risiko kebijakan secara bersyarat. Model belum layak dipakai untuk meramalkan persentase penggunaan AI atau dampak numerik di sekolah nyata karena belum dikalibrasi dengan data siswa."
    )
    for run in p.runs[1:]:
        set_run_font(run, size=10.6)
    doc.add_page_break()
    # Executive summary
    add_heading(doc, "Ringkasan eksekutif", 1)
    add_body(
        doc,
        "Kami membangun kelas sintetis berisi 30 siswa. Setiap siswa memilih cara menyelesaikan tugas: mengerjakan sendiri, berdiskusi, meminta jawaban dari teman, atau menggunakan AI ketika kebijakan menyediakan AI. Pilihan berikutnya dipengaruhi oleh hasil yang dialami siswa sendiri. Tidak ada tipe tetap seperti siswa rajin, penyalin, pengguna AI, atau broker. Karena itu, pola penggunaan, sumber jawaban, dan penggantian interaksi manusia dapat muncul dari interaksi lokal.",
    )
    add_body(
        doc,
        "Eksperimen E8 menjalankan 23.000 simulasi. Perbandingan utama memakai 400 kelas berpasangan untuk setiap kebijakan. Ketiga versi kelas menerima siswa, jaringan, soal, dan kejadian acak yang sepadan; perbedaannya adalah layanan AI. Cara ini membuat perbandingan internal lebih bersih daripada membandingkan tiga kelas acak yang tidak setara.",
    )
    add_body(
        doc,
        "Dalam model, AI jawaban langsung lebih sering dipakai dan menaikkan nilai tugas, tetapi menurunkan pengetahuan dan ujian tanpa AI. Tutor terkendali lebih jarang dipakai, namun meningkatkan pengetahuan dan ujian. Hasil ini searah dengan eksperimen acak yang relevan, tetapi besarnya efek tidak boleh dibaca sebagai prediksi dunia nyata [1, 2].",
    )
    add_table(
        doc,
        ["Tingkat keyakinan", "Penilaian", "Arti praktis"],
        [
            ["Integritas komputasi", "Tinggi", "Desain dikunci, keluaran diaudit, dan 222 tes lulus."],
            ["Perbandingan kausal di dalam model", "Cukup tinggi", "Kebijakan berbeda, tetapi dunia acak dipasangkan pada persamaan yang sama."],
            ["Arah AI langsung dibanding tutor", "Sedang", "Arah bertahan pada aturan keputusan dan horizon yang diuji serta sesuai dengan beberapa studi."],
            ["Angka penggunaan dan besar dampak", "Rendah", "Belum dikalibrasi; aturan keputusan mengubah penggunaan AI langsung dari sekitar 31 persen menjadi 81 persen."],
            ["Ramalan untuk sekolah tertentu", "Rendah", "Tidak ada data sekolah, dialog AI nyata, adaptasi guru, atau perilaku melewati pagar."],
        ],
        [1.55, 1.15, 4.15],
    )

    add_heading(doc, "Cara membaca laporan", 1)
    add_body(doc, "Laporan ini membedakan tiga lapisan agar pembaca tidak mencampur asumsi dengan hasil.")
    add_numbered(
        doc,
        [
            "Aturan yang dipasang adalah mekanisme yang sengaja ditulis ke dalam model, misalnya kapasitas berbagi dua jawaban per tugas atau AI tutor yang menahan jawaban akhir.",
            "Perilaku yang muncul adalah pola yang tidak ditentukan per siswa, misalnya siapa yang menjadi sumber jawaban, berapa banyak siswa terus memakai AI, atau apakah interaksi manusia berkurang.",
            "Bukti dari literatur menjelaskan apakah arah suatu mekanisme masuk akal. Literatur tidak otomatis membenarkan angka parameter kita.",
        ],
    )
    add_body(doc, "Istilah data dalam laporan berarti data sintetis hasil Monte Carlo. Data ini bukan rekam jejak siswa nyata dan tidak mengandung data pribadi.")

    add_heading(doc, "Pertanyaan penelitian", 1)
    add_body(
        doc,
        "Pertanyaan pokoknya adalah bagaimana kebijakan AI mengubah pilihan siswa, hasil belajar, dan jaringan bantuan ketika semua siswa memiliki akses tetapi bebas memilih untuk menggunakan atau tidak menggunakan AI.",
    )
    add_bullets(
        doc,
        [
            "Apakah penggunaan AI muncul tanpa bonus sosial khusus untuk AI dan tanpa tipe pengguna yang ditetapkan sebelumnya?",
            "Apakah AI jawaban langsung memisahkan nilai tugas dari kemampuan tanpa bantuan?",
            "Apakah tutor terkendali dapat menjaga atau meningkatkan belajar meskipun penggunaannya menuntut usaha?",
            "Apakah kedua layanan menggantikan diskusi dan bantuan manusia?",
            "Seberapa sensitif kesimpulan terhadap aturan memilih, kekuatan belajar, biaya usaha, bentuk jaringan, dan lamanya semester?",
        ],
    )

    add_heading(doc, "Apa yang dimaksud model berbasis agen", 1)
    add_body(
        doc,
        "Model berbasis agen adalah kelas virtual yang berisi individu sederhana. Setiap individu memiliki keadaan sendiri, melihat informasi yang terbatas, membuat pilihan, menerima hasil, lalu memperbarui pengalamannya. Peneliti mengulang kelas virtual berkali-kali untuk melihat pola yang sering muncul. Tujuan utama model seperti ini dapat berupa penjelasan mekanisme dan pengujian skenario, bukan selalu ramalan [13, 14].",
    )
    add_body(
        doc,
        "Analogi yang dekat dengan pekerjaan guru adalah laboratorium pemikiran. Guru dapat menanyakan apa yang terjadi jika jawaban AI sempurna tetapi proses belajar sangat kecil, atau jika tutor memberi petunjuk yang menuntut usaha. Model menjawab menurut aturan yang ditulis. Jika aturan tidak realistis, jawaban model juga tidak boleh dibawa langsung ke kebijakan nyata.",
    )

    add_heading(doc, "Bagian model satu per satu", 1)
    add_heading(doc, "Simpul siswa", 2)
    add_body(doc, "Satu simpul atau lingkaran mewakili satu siswa sintetis. E8 menggunakan 30 siswa per kelas. Identitas simpul tidak merujuk kepada siswa nyata.")
    add_table(
        doc,
        ["Keadaan siswa", "Makna sederhana", "Peran dalam keputusan"],
        [
            ["K pengetahuan", "Peluang dasar menjawab benar tanpa bantuan", "Mempengaruhi hasil tugas dan ujian"],
            ["C catatan kepercayaan diri", "Rata-rata bergerak dari nilai tugas", "Tidak memengaruhi pilihan; bukan ukuran psikologis tervalidasi"],
            ["alpha efisiensi belajar", "Seberapa banyak pengetahuan bertambah setelah belajar", "Membuat siswa berbeda secara acak"],
            ["kappa biaya usaha", "Seberapa berat usaha dirasakan", "Mengurangi imbalan tindakan yang memerlukan usaha"],
            ["Q nilai tindakan", "Catatan pribadi tentang seberapa menguntungkan suatu pilihan", "Menentukan pilihan berikutnya bersama eksplorasi"],
            ["Aij keyakinan sumber", "Perkiraan siswa i bahwa siswa j dapat dan mau memberi jawaban", "Menentukan siapa yang diminta; hanya berubah setelah interaksi langsung"],
        ],
        [1.45, 2.65, 2.75],
    )

    add_heading(doc, "Hubungan sosial", 2)
    add_body(
        doc,
        "Garis antara dua siswa berarti mereka memiliki peluang kontak yang stabil sehingga dapat berdiskusi atau saling meminta jawaban. Garis tidak berarti persahabatan pasti, kedekatan emosional, atau pengaruh sebab akibat. Pada perbandingan utama E8, semua garis memiliki bobot yang sama. Kerangka mendukung bobot kontak, tetapi E8 tidak menggunakan bobot yang diukur dari kelas nyata.",
    )
    add_body(
        doc,
        "Jaringan dasar dibuat dengan stochastic block model. Metode ini menghasilkan kelompok yang lebih rapat di dalam kelompok daripada antar kelompok [12]. Pilihan ini cocok sebagai kelas hipotetis yang mempunyai lingkaran pertemanan, tetapi bukan bukti bahwa semua kelas mengikuti bentuk yang sama.",
    )
    add_figure(
        doc,
        figures["anatomy"],
        "Gambar 1. Komponen jaringan pada satu kelas sintetis yang representatif.",
        "Tiga kelompok siswa terhubung oleh garis kontak sosial. Simpul lebih besar memiliki pengetahuan awal lebih tinggi. Sebuah panah menunjukkan bahwa keyakinan sumber dan aliran jawaban memiliki arah. AI berada di luar jaringan siswa.",
    )

    add_heading(doc, "Tiga jaringan yang tidak boleh dicampur", 2)
    add_table(
        doc,
        ["Lapisan", "Bentuk", "Fungsi", "Apakah memengaruhi masa depan"],
        [
            ["Jaringan kontak G", "Garis dua arah dan tetap", "Menentukan siapa dapat berinteraksi", "Ya sebagai peluang kontak"],
            ["Keyakinan pribadi Aij", "Nilai berarah milik peminta", "Menilai keandalan calon sumber", "Ya, tetapi hanya bagi siswa yang mengalami interaksi"],
            ["Jaringan aliran H", "Panah transfer jawaban yang terjadi", "Merekam rute dan sumber yang muncul", "Tidak; hanya keluaran analisis"],
        ],
        [1.35, 1.55, 2.25, 1.55],
    )
    add_body(
        doc,
        "Pemisahan ini melindungi klaim emergence. Siswa tidak menjadi sumber karena model memberi label sumber. Ia menjadi sumber karena memiliki jawaban, dipilih, bersedia berbagi, dan tidak kehabisan kapasitas. Setelah itu peneliti menyebutnya sumber berdasarkan aliran yang tercatat.",
    )

    add_heading(doc, "Pilihan tindakan", 2)
    add_table(
        doc,
        ["Kode", "Tindakan", "Apa yang terjadi"],
        [
            ["S", "Mengerjakan sendiri", "Siswa berusaha menjawab dan belajar dari pengerjaan."],
            ["D", "Diskusi", "Siswa memilih tetangga sosial dan memperoleh proses belajar bersama yang disederhanakan."],
            ["C", "Meminta jawaban teman", "Siswa memilih calon sumber menurut kontak dan keyakinan pribadi. Keberhasilan tergantung kepemilikan jawaban, kemauan berbagi, dan kapasitas."],
            ["A", "Menggunakan AI", "Tersedia mulai tugas kedua pada dua kebijakan AI. Semua siswa memiliki akses, tetapi penggunaan tetap pilihan."],
        ],
        [0.55, 1.65, 4.5],
    )

    add_heading(doc, "Urutan satu tugas", 2)
    add_numbered(
        doc,
        [
            "Model membekukan keadaan siswa, jaringan, dan keyakinan sebelum tugas dimulai.",
            "Setiap siswa memilih tindakan. Pada tugas pertama terdapat benih penyalinan sekali saja; mulai tugas kedua ada eksplorasi netral enam persen.",
            "Siswa yang berdiskusi atau meminta jawaban memilih tetangga berdasarkan informasi lokal.",
            "Permintaan jawaban diproses serentak per putaran. Jawaban yang baru diterima pada satu putaran baru dapat diteruskan pada putaran berikutnya.",
            "Setiap siswa menerima nilai tugas, biaya usaha, dan perubahan pengetahuan sesuai tindakan dan hasilnya.",
            "Siswa memperbarui nilai tindakan yang baru dipilih. Keyakinan terhadap sumber berubah hanya setelah permintaan langsung.",
            "Catatan kepercayaan diri bergerak mengikuti nilai tugas, tetapi tidak mengubah pilihan.",
        ],
    )

    add_heading(doc, "Bagaimana siswa memilih", 2)
    add_body(
        doc,
        "Pada tugas pertama, sepuluh persen peluang diberikan kepada penyalinan sebagai benih awal. Benih ini dimatikan setelah tugas pertama. Tanpa kemungkinan awal, model tidak dapat menciptakan tindakan dari ketiadaan. Mulai tugas kedua, siswa biasanya memilih tindakan dengan nilai pribadi tertinggi. Dengan peluang enam persen, siswa mengeksplorasi tindakan yang tersedia secara netral. AI tidak mendapat bonus karena populer, dan pilihan teman tidak langsung menaikkan daya tarik AI.",
    )
    add_body(
        doc,
        "Nilai tindakan diperbarui dengan aturan sederhana: perkiraan lama digeser ke arah imbalan yang baru dialami. Ini mengambil ide umum dari pembelajaran penguatan dan Q learning [11]. Namun, model tidak membuktikan bahwa siswa manusia menghitung Q secara sadar atau mengikuti persamaan yang sama.",
    )

    add_heading(doc, "Bagaimana meminta dan berbagi jawaban", 2)
    add_body(
        doc,
        "Seorang peminta lebih cenderung memilih tetangga yang sebelumnya dinilai dapat dan mau memberi jawaban. Pemilik jawaban memutuskan apakah akan berbagi. Kesediaan dapat meningkat bila ia sendiri pernah berhasil menerima bantuan dari peminta. Setiap pemilik dapat memasok paling banyak dua transfer per tugas. Literatur tentang resiprositas mendukung kemungkinan bahwa pengalaman bantuan memengaruhi kerja sama [10], tetapi nilai dua transfer dan kekuatan resiprositas dalam model adalah pilihan skenario, bukan hasil pengukuran kelas.",
    )

    add_heading(doc, "Tiga kebijakan AI", 2)
    add_table(
        doc,
        ["Kebijakan", "Respons layanan", "Belajar yang dipasang", "Biaya usaha"],
        [
            ["Tanpa AI", "Hanya S D dan C", "Sesuai mekanisme manusia V3", "Sesuai tindakan"],
            ["AI jawaban langsung", "Jawaban lengkap dengan ketepatan 100 persen pada skenario acuan", "Sepuluh persen dari kenaikan belajar tindakan menyalin", "Nol pada skenario acuan"],
            ["AI tutor terkendali", "Petunjuk bertahap, jawaban akhir ditahan, usaha diwajibkan", "Setara kenaikan belajar mengerjakan sendiri pada skenario acuan", "Satu"],
        ],
        [1.35, 2.45, 2.25, 0.8],
    )
    add_body(
        doc,
        "AI adalah layanan eksternal, bukan simpul yang terhubung ke semua siswa. Membuat AI sebagai simpul pusat akan mencampur akses teknologi dengan struktur pertemanan dan secara otomatis menjadikannya broker terbesar. E8 sengaja memisahkan keduanya.",
    )

    add_heading(doc, "Nilai tugas pengetahuan ujian dan kepercayaan diri", 2)
    add_body(
        doc,
        "Nilai tugas mengukur apakah jawaban saat bantuan tersedia benar. Pengetahuan adalah keadaan internal sintetis yang memengaruhi peluang menjawab sendiri. Ujian akhir dikerjakan tanpa AI dan tidak mengubah keadaan siswa. Selisih nilai tugas dan ujian dipakai sebagai indikator kinerja yang tampak lebih tinggi daripada kemampuan tanpa bantuan. Eksperimen Bastani dan kolega menunjukkan mengapa pemisahan bantuan saat latihan dari ujian tanpa bantuan relevan [1]. Praktik pengambilan kembali juga menunjukkan bahwa pengujian dapat mengukur sekaligus mendukung retensi, walau ujian E8 sengaja dibuat hanya sebagai alat ukur [7].",
    )
    add_body(
        doc,
        "Variabel C tidak boleh ditafsirkan sebagai rasa percaya diri psikologis. Ia hanya rata-rata bergerak dari nilai tugas dan tidak memengaruhi perilaku. Nama yang lebih aman untuk komunikasi eksternal adalah catatan keberhasilan yang dirasakan. Tanpa instrumen survei yang tervalidasi, hasil C tidak dapat dipakai untuk menyimpulkan kesehatan mental, keyakinan akademik, atau ketergantungan psikologis.",
    )

    add_heading(doc, "Apa yang dipasang dan apa yang muncul", 1)
    add_table(
        doc,
        ["Komponen", "Status", "Penjelasan"],
        [
            ["Tidak ada bonus prevalensi AI atau menyalin", "Dipasang", "Siswa tidak menyukai tindakan hanya karena banyak tetangga melakukannya."],
            ["Benih penyalinan sepuluh persen pada tugas pertama", "Dipasang", "Memastikan tindakan dapat muncul, lalu dimatikan."],
            ["Eksplorasi netral enam persen", "Dipasang", "Memberi kesempatan mencoba semua tindakan yang tersedia."],
            ["Distribusi pengetahuan efisiensi dan biaya", "Dipasang", "Heterogenitas diacak dan tidak diikat ke posisi jaringan."],
            ["Jawaban AI langsung sempurna dan murah", "Dipasang pada acuan", "Skenario batas yang sengaja menarik, bukan estimasi layanan nyata."],
            ["Tutor menahan jawaban dan mendukung belajar", "Dipasang pada acuan", "Mekanisme diuji pada rentang parameter, bukan ditemukan oleh model."],
            ["Siapa memakai AI dan berapa lama", "Muncul", "Ditentukan oleh percobaan netral dan pengalaman imbalan."],
            ["Siapa menjadi sumber jawaban", "Muncul", "Tidak ada label sumber atau broker."],
            ["Konsentrasi aliran bantuan", "Muncul", "Dihitung setelah simulasi dari transfer yang benar-benar terjadi."],
            ["Penggantian diskusi dan bantuan manusia", "Muncul", "Perubahan komposisi tindakan akibat persaingan antar pilihan."],
            ["Selisih nilai tugas dan ujian", "Muncul secara bersyarat", "Muncul dari mekanisme belajar dan bantuan yang dipasang."],
        ],
        [2.45, 1.2, 3.0],
        font_size=8.7,
    )

    add_heading(doc, "Metodologi eksperimen E8", 1)
    add_heading(doc, "Desain berpasangan", 2)
    add_body(
        doc,
        "Dalam satu replikasi, kondisi tanpa AI, AI jawaban langsung, dan tutor terkendali memakai jaringan, sifat siswa, soal, dan aliran acak yang sepadan. Perbandingan itu menyerupai tiga dunia paralel yang berbeda hanya pada kebijakan AI. Teknik common random numbers mengurangi gangguan acak ketika dua kebijakan dibandingkan.",
    )
    add_heading(doc, "Jumlah simulasi", 2)
    add_table(
        doc,
        ["Blok", "Jumlah run", "Tujuan"],
        [
            ["Perbandingan utama", "1.200", "400 replikasi dikali tiga kebijakan acuan"],
            ["Permukaan mekanisme", "7.400", "Menguji ketepatan AI langsung, pembelajaran insidental, kekuatan tutor, scaffolding, dan biaya"],
            ["Ketahanan", "14.400", "Menguji aturan keputusan, pembelajaran sosial, topologi, pengacakan sifat, dan horizon"],
            ["Total", "23.000", "Seluruh run E8 yang dilaporkan"],
        ],
        [1.75, 1.15, 3.75],
    )
    add_heading(doc, "Praregistrasi dan pemeriksaan", 2)
    add_body(
        doc,
        "Pertanyaan, prediksi, ambang praktis, dan klasifikasi kebijakan dikunci sebelum hasil dilihat. Mesin V3 tidak diubah selama E8. Audit menguji 32.615 berkas tanpa berkas hilang atau hash tidak cocok. Seluruh 222 tes perangkat lunak lulus. Pemeriksaan ini mendukung integritas komputasi, tetapi tidak menggantikan validasi terhadap siswa nyata.",
    )
    add_heading(doc, "Monte Carlo dan interval", 2)
    add_body(
        doc,
        "Setiap run adalah satu kelas hipotetis dengan variasi acak. Rata-rata 400 kelas mengurangi pengaruh kebetulan dari satu kelas. Interval ketidakpastian pada laporan E8 menunjukkan variasi Monte Carlo di dalam model. Interval itu bukan interval populasi siswa Indonesia karena model tidak mengambil sampel dari populasi tersebut.",
    )

    add_heading(doc, "Hasil utama", 1)
    add_figure(
        doc,
        figures["primary"],
        "Gambar 2. Pilihan dan hasil rata-rata pada perbandingan utama E8.",
        "Grafik kiri membandingkan persentase tindakan akhir. Grafik kanan membandingkan nilai tugas, pengetahuan, ujian tanpa AI, dan selisih nilai tugas dengan ujian pada tiga kebijakan.",
    )
    primary_rows = []
    labels = [
        ("pA_late", "Penggunaan AI akhir"),
        ("pS_late", "Mengerjakan sendiri"),
        ("pD_late", "Diskusi"),
        ("pC_late", "Menyalin dari teman"),
        ("grade_late", "Nilai tugas akhir"),
        ("K_final", "Pengetahuan akhir"),
        ("exam_final", "Ujian tanpa AI"),
        ("grade_exam_gap", "Selisih nilai tugas dan ujian"),
    ]
    for key, label in labels:
        primary_rows.append([label] + [f"{means[c][key] * 100:.1f}" for c in REGIME_ORDER])
    add_table(doc, ["Ukuran dalam poin", "Tanpa AI", "AI langsung", "Tutor"], primary_rows, [2.7, 1.25, 1.25, 1.25])
    add_body(
        doc,
        "Dibanding tanpa AI, AI jawaban langsung menaikkan nilai tugas 4,8 poin, menurunkan pengetahuan 3,3 poin, menurunkan ujian tanpa AI 3,4 poin, dan memperlebar selisih nilai tugas dengan ujian 8,2 poin. Tutor terkendali menaikkan nilai tugas 1,0 poin, menaikkan pengetahuan 2,9 poin, menaikkan ujian 2,9 poin, dan memperkecil selisih 1,9 poin.",
    )
    add_body(
        doc,
        "Perbandingan tutor dengan AI langsung lebih stabil daripada angka absolutnya. Tutor unggul sekitar 6,2 poin pada pengetahuan dan ujian. Arah ini tetap pada argmax, dua softmax, dan horizon 20 maupun 40 tugas. Namun, penggunaan AI langsung berubah besar menurut aturan keputusan, sehingga persentase adopsi tidak dapat dipercaya sebagai ramalan.",
    )
    add_figure(
        doc,
        figures["trajectory"],
        "Gambar 3. Perjalanan penggunaan AI, pengetahuan, dan nilai tugas selama 20 tugas.",
        "Tiga grafik garis menunjukkan rata-rata penggunaan AI, pengetahuan, dan nilai tugas dari tugas pertama sampai tugas kedua puluh untuk tiga kebijakan.",
    )

    add_heading(doc, "Cara membaca grafik jaringan hasil", 1)
    add_body(
        doc,
        "Gambar berikut menggunakan replikasi 282. Replikasi ini dipilih sebelum visualisasi dengan menghitung jarak baku terhadap rata-rata 400 kelas pada sepuluh ukuran utama di ketiga kebijakan. Ia adalah kelas yang paling dekat dengan pusat hasil, bukan kelas dengan gambar paling dramatis. Karena tetap hanya satu replikasi, angka kebijakan harus dibaca dari agregat 400 kelas, bukan dari panel ini.",
    )
    add_figure(
        doc,
        figures["flows"],
        "Gambar 4. Jaringan bantuan manusia pada lima tugas terakhir dalam tiga dunia berpasangan.",
        "Tiga panel memakai posisi siswa yang sama. Garis abu-abu adalah kontak sosial. Panah biru adalah transfer jawaban manusia yang berhasil. Ukuran simpul menunjukkan jumlah jawaban yang disuplai. Warna ungu menunjukkan proporsi penggunaan AI oleh siswa.",
    )
    add_bullets(
        doc,
        [
            "Garis abu-abu yang sama di semua panel menegaskan bahwa kebijakan tidak mengubah jaringan kontak dasar.",
            "Panah biru berkurang ketika AI tersedia karena sebagian siswa mengganti bantuan manusia dengan layanan eksternal.",
            "Simpul besar menunjukkan sumber yang muncul dari kejadian transfer, bukan siswa yang ditandai sebagai sumber sejak awal.",
            "Ungu menunjukkan penggunaan AI per siswa. AI tidak digambar sebagai simpul pusat karena akses bukan hubungan pertemanan.",
        ],
    )
    add_body(
        doc,
        "Pada agregat 400 kelas, bagian jawaban yang disuplai oleh 20 persen sumber teratas meningkat dari 62,5 persen tanpa AI menjadi 81,8 persen pada AI langsung dan 75,6 persen pada tutor. Gini sumber juga naik dari 0,647 menjadi 0,772 dan 0,736. Artinya, bantuan manusia yang tersisa menjadi lebih terkonsentrasi dalam model.",
    )

    add_heading(doc, "Penyebaran pengaruh", 1)
    add_heading(doc, "Penyebaran tanpa pembelajaran sosial", 2)
    add_body(
        doc,
        "Karena semua siswa memiliki akses, penyebaran bukan penyebaran akses. Pada model dasar, siswa mencoba AI melalui eksplorasi netral lalu mempertahankan atau meninggalkannya menurut pengalaman sendiri. Assortativitas adopsi dan hubungan dengan penggunaan tetangga sedikit negatif. Model dasar tidak menunjukkan penularan penggunaan AI antar teman.",
    )
    add_heading(doc, "Penyebaran dengan pembelajaran sosial generik", 2)
    add_body(
        doc,
        "Dalam uji ketahanan opsional, siswa mengamati tindakan dan imbalan tetangga untuk semua tindakan secara simetris. Aturan ini tidak memberi bonus khusus kepada AI. Karena jawaban AI langsung sangat menguntungkan dalam jangka pendek, pengamatan sosial menaikkan penggunaan AI langsung sekitar 25,9 poin dan menurunkan ujian sekitar 5,5 poin. Pada tutor, penggunaan naik sekitar 12,9 poin dan ujian naik tipis 0,4 poin. Ini adalah konsekuensi aturan model, bukan bukti laju penularan AI di kelas nyata. Literatur menunjukkan pembelajaran sosial dapat menguntungkan, tetapi strategi menyalin bergantung pada informasi dan lingkungan [9].",
    )
    add_body(
        doc,
        "Jika data jaringan nyata kelak tersedia, korelasi antar teman tetap tidak cukup untuk membuktikan pengaruh. Homofili, pemilihan teman, dan pengaruh sosial dapat menghasilkan pola yang sama. Literatur metodologis menunjukkan ketiganya sulit dipisahkan tanpa asumsi atau desain yang kuat [15].",
    )

    add_heading(doc, "Parameter yang paling memengaruhi hasil", 1)
    add_table(
        doc,
        ["Komponen", "Temuan sensitivitas E8", "Implikasi"],
        [
            ["Aturan keputusan", "Penggunaan AI langsung sekitar 31 persen dengan argmax, 81 persen dengan softmax 0,10, dan 56 persen dengan softmax 0,25", "Pendorong terbesar untuk angka adopsi. Angka adopsi tidak defensibel tanpa data perilaku."],
            ["Kekuatan belajar tutor", "Menaikkan mT dari 0,50 ke 1,50 meningkatkan pengetahuan sekitar 3,17 poin", "Pendorong utama hasil belajar tutor."],
            ["Biaya usaha tutor", "Biaya dua dibanding satu menurunkan penggunaan sekitar 10,52 poin", "Pagar yang terlalu berat dapat efektif namun diabaikan."],
            ["Pembelajaran sosial", "Menaikkan penggunaan AI langsung sekitar 25,9 poin dan memperbesar kerugian ujian", "Mekanisme pengaruh teman penting bila benar-benar terjadi."],
            ["Topologi SBM atau ER", "Dampak marjinal hampir nol pada adopsi dan hasil utama", "Wajar karena semua siswa langsung memperoleh akses AI; jaringan tetap penting untuk bantuan manusia."],
            ["Horizon 20 atau 40", "Arah kerugian AI langsung dan keuntungan tutor bertahan", "Hasil utama bukan hanya kejadian singkat di akhir 20 tugas."],
        ],
        [1.55, 2.75, 2.7],
        font_size=8.7,
    )

    add_heading(doc, "Dasar ilmiah dan batas setiap klaim", 1)
    add_body(
        doc,
        "Tabel berikut menjawab pertanyaan paling ketat: sumber apa yang mendukung asumsi kita, dan bagian mana yang tetap buatan model. Tidak semua parameter boleh diberi legitimasi dari literatur umum.",
    )
    evidence_rows = [
        ["AI tanpa pagar dapat memisahkan kinerja terbantu dari ujian tanpa bantuan", "[1] eksperimen acak hampir 1.000 siswa matematika sekolah menengah", "Mendukung arah risiko. Tidak membenarkan akurasi 100 persen, biaya nol, mU 0,10, atau besar efek E8."],
        ["Tutor dengan petunjuk, umpan balik, dan usaha dapat mendukung belajar", "[1] [2] [3] [4] [8]", "Mendukung mekanisme umum. Tidak membuktikan semua tutor atau nilai mT dan scaffolding E8 efektif."],
        ["Diskusi dan aktivitas aktif dapat membantu belajar", "[5] meta-analisis 225 studi STEM", "Mendukung adanya jalur belajar dari aktivitas dan diskusi. Tidak mengkalibrasi persamaan D atau alpha."],
        ["Ujian tanpa bantuan memberi informasi berbeda dari nilai terbantu", "[1] serta literatur praktik pengambilan kembali [7]", "Mendukung pemisahan pengukuran. Jumlah 30 soal dan peluang benar di E8 tetap buatan."],
        ["Imbalan pengalaman dapat memperbarui nilai tindakan", "[11] metode Q learning", "Mendukung algoritme komputasi. Bukan bukti bahwa siswa menggunakan pembaruan Q yang sama."],
        ["Pengamatan teman dapat memengaruhi strategi", "[9] turnamen strategi pembelajaran sosial", "Mendukung kemungkinan umum. Tidak memberi angka eta sosial atau laju adopsi AI siswa."],
        ["Resiprositas dapat menopang kerja sama", "[10] tinjauan resiprositas", "Mendukung arah kemauan berbagi. Beta resiprositas dan kapasitas dua tidak dikalibrasi."],
        ["Kelompok sosial dapat direpresentasikan dengan blok jaringan", "[12] stochastic block model dan [16] homofili", "Mendukung metode pembentukan kelompok. Rasio blok, derajat enam, dan jaringan statis bukan fakta kelas tertentu."],
        ["Interaksi kelas dapat dianalisis sebagai jaringan", "[17] studi jaringan interaksi dan nilai fisika", "Mendukung kegunaan jaringan. Hubungan yang dilaporkan bersifat kontekstual dan tidak membuktikan kausalitas E8."],
        ["Kepercayaan pribadi terhadap sumber berubah hanya dari pengalaman langsung", "Tidak ada validasi langsung untuk persamaan Aij", "Asumsi transparan untuk mencegah reputasi global buatan. Harus diuji dengan data permintaan bantuan."],
        ["Siswa berbeda dalam K0 alpha dan kappa tetapi sifat independen dari posisi", "Pilihan desain anti tipe", "Menjaga agar broker dan pengguna tidak ditetapkan. Bentuk distribusi dan independensinya belum empiris."],
        ["C mengikuti nilai tugas", "Tidak diperlakukan sebagai klaim psikologis", "Hanya alat ukur internal. Jangan menyebutnya kepercayaan diri nyata tanpa instrumen yang tervalidasi."],
    ]
    add_table(doc, ["Asumsi atau klaim", "Bukti sejawat", "Yang belum didukung"], evidence_rows, [2.15, 2.0, 2.85], font_size=8.1)

    add_heading(doc, "Keterbatasan", 1)
    add_heading(doc, "Tidak ada kalibrasi data nyata", 2)
    add_body(
        doc,
        "Ini adalah keterbatasan terbesar. Model tidak mengetahui seberapa sering siswa Indonesia memakai AI, seberapa banyak mereka belajar dari satu interaksi, berapa biaya usaha yang dirasakan, atau bagaimana jaringan bantuan terbentuk di sekolah tertentu. Karena itu, 31 persen penggunaan AI langsung dan kenaikan 2,9 poin pada tutor adalah keluaran skenario, bukan perkiraan lapangan.",
    )
    add_heading(doc, "Jaringan sosial masih statis dan tidak berbobot secara empiris", 2)
    add_body(
        doc,
        "Kontak dasar tidak berubah selama semester. Model dapat mencatat jaringan bantuan yang muncul, tetapi jaringan itu tidak mengubah hubungan kontak pada masa depan. Bobot tepi pada perbandingan utama sama dengan satu. Model belum menangkap persahabatan yang menguat, konflik, perubahan kelompok, atau saluran berbeda untuk diskusi dan penyalinan.",
    )
    add_heading(doc, "Tindakan dan pembelajaran disederhanakan", 2)
    add_body(
        doc,
        "Satu tindakan merangkum proses yang dalam kenyataan dapat berlangsung lama. Diskusi tidak memiliki isi percakapan. Penyalinan tidak mempunyai deteksi guru. Tutor tidak berdialog, menilai miskonsepsi, atau menghadapi prompt untuk melewati pagar. Hasil belajar dihitung segera setelah tindakan, bukan melalui memori, lupa, transfer lintas topik, atau retensi jangka panjang.",
    )
    add_heading(doc, "AI acuan terlalu bersih", 2)
    add_body(
        doc,
        "AI jawaban langsung acuan selalu benar dan gratis dalam usaha. Tutor mempunyai solusi terverifikasi serta umpan balik yang tepat. Layanan nyata dapat salah, berubah kualitas, memakai model berbeda, atau digunakan di luar antarmuka sekolah. E8 menguji konsekuensi dua desain layanan ideal, bukan seluruh ekosistem AI.",
    )
    add_heading(doc, "Guru dan sekolah tidak beradaptasi", 2)
    add_body(
        doc,
        "Guru tidak mengubah tugas, penilaian, kelompok, atau aturan setelah melihat perilaku siswa. Tidak ada kebijakan privasi, disiplin, pemeriksaan plagiarisme, biaya perangkat, kesenjangan akses rumah, maupun pelatihan literasi AI. Padahal respons sekolah dapat mengubah perilaku sama kuatnya dengan teknologi.",
    )
    add_heading(doc, "Generalisasi konteks terbatas", 2)
    add_body(
        doc,
        "Bukti eksperimen utama berasal dari matematika sekolah menengah di satu sekolah dan fisika perguruan tinggi. Tinjauan ITS K sampai 12 menemukan hasil yang umumnya positif tetapi heterogen [4]. Kesimpulan tidak boleh dipindahkan otomatis ke bahasa, seni, pendidikan vokasi, atau kebutuhan khusus.",
    )
    add_heading(doc, "Pengaruh sosial sulit divalidasi", 2)
    add_body(
        doc,
        "Jika dua teman sama-sama memakai AI, penyebabnya dapat berupa pengaruh langsung, kesamaan minat, guru yang sama, atau akses yang sama. Model dapat memisahkan mekanisme karena peneliti mengaturnya, tetapi data observasional nyata tidak otomatis dapat melakukan hal yang sama [15].",
    )

    add_heading(doc, "Data sintetis ini paling tepat mewakili apa", 1)
    add_table(
        doc,
        ["Layak diwakili", "Tidak layak diwakili"],
        [
            ["Konsekuensi logis dari mekanisme yang tertulis", "Persentase pengguna AI di sekolah tertentu"],
            ["Arah tradeoff antara imbalan langsung, usaha, dan belajar", "Besar efek kausal dunia nyata"],
            ["Perbandingan bersyarat antara desain AI pada dunia acak yang sama", "Perilaku individu siswa atau diagnosis siswa"],
            ["Pola yang perlu dicari dalam data nyata seperti selisih nilai dan konsentrasi bantuan", "Prediksi nilai, kepercayaan diri, atau risiko menyontek seseorang"],
            ["Uji apakah kesimpulan bertahan ketika asumsi diubah", "Bukti bahwa pengaruh teman benar-benar terjadi"],
            ["Panduan memilih data yang perlu dikumpulkan", "Pembenaran otomatis untuk melarang atau mewajibkan AI"],
        ],
        [3.35, 3.35],
    )
    add_body(
        doc,
        "Kalimat yang paling aman adalah: jika siswa belajar dari imbalan pribadi seperti dalam model, jika AI langsung menggantikan usaha belajar, dan jika tutor menghasilkan belajar setara mengerjakan sendiri, maka model memprediksi arah hasil yang dilaporkan. Menghapus kata jika akan mengubah hasil bersyarat menjadi klaim dunia nyata yang tidak didukung.",
    )

    add_heading(doc, "Kemungkinan penggunaan", 1)
    add_table(
        doc,
        ["Penggunaan", "Cara yang bertanggung jawab", "Batas"],
        [
            ["Laboratorium kebijakan", "Bandingkan pagar AI, biaya usaha, dan kekuatan belajar sebelum uji lapangan", "Gunakan rentang, bukan satu prediksi"],
            ["Perancangan pengajaran", "Jaga diskusi manusia, tugas tanpa AI, dan petunjuk yang menuntut usaha", "Model tidak menentukan metode terbaik untuk semua mata pelajaran"],
            ["Perancangan pilot", "Pilih indikator seperti penggunaan, pengabaian, nilai terbantu, ujian tanpa AI, dan aliran bantuan", "Pilot memerlukan persetujuan etika dan perlindungan privasi"],
            ["Sistem peringatan tingkat kelas", "Pantau kenaikan selisih nilai tugas dan ujian atau hilangnya diskusi secara agregat", "Jangan menandai atau menghukum siswa berdasarkan model"],
            ["Uji sistem agentik", "Gunakan V3 sebagai pembanding transparan untuk memeriksa apakah agen bahasa mengubah mekanisme", "Agen bahasa tidak boleh menggantikan validasi lapangan"],
        ],
        [1.55, 3.05, 2.2],
        font_size=8.6,
    )

    add_heading(doc, "Apa yang perlu dilakukan sebelum sistem lebih agentik", 1)
    add_numbered(
        doc,
        [
            "Bekukan V3 dan E8 sebagai baseline transparan. Jangan mengubahnya ketika menambahkan komponen bahasa.",
            "Pertahankan keputusan siswa yang dapat diaudit. Jika model bahasa ditambahkan, mulai dari mengganti kualitas respons tutor, bukan memberi model bahasa kendali penuh atas pilihan siswa.",
            "Uji apakah komponen agentik mereproduksi amplop mekanisme E8 pada kondisi sederhana. Jika tidak, identifikasi aturan baru yang menyebabkan perbedaan.",
            "Catat setiap prompt, respons, kesalahan, petunjuk, percobaan siswa, dan penggunaan jawaban akhir. Tanpa log ini, alasan hasil akan kembali menjadi kotak hitam.",
            "Jalankan ulang dengan beberapa model bahasa dan versi model. Hasil yang hanya muncul pada satu model atau satu prompt tidak cukup stabil untuk rekomendasi sekolah.",
            "Pisahkan uji teknis dari klaim pendidikan. Agen dapat membuat interaksi lebih realistis, tetapi tidak otomatis membuat parameter siswa lebih benar.",
        ],
    )

    add_heading(doc, "Data minimum untuk validasi masa depan", 1)
    add_body(
        doc,
        "Jika sekolah belum dapat menyediakan data sekarang, model masih dapat dipakai sebagai alat perumusan hipotesis. Ketika kesempatan pilot muncul, kumpulkan data minimum berikut secara agregat atau dipseudonimkan.",
    )
    add_table(
        doc,
        ["Data", "Mengapa diperlukan", "Contoh pengukuran aman"],
        [
            ["Penggunaan AI per tugas", "Mengkalibrasi percobaan, persistensi, dan pengabaian", "Log mode AI tanpa isi percakapan untuk analisis awal"],
            ["Nilai dengan bantuan dan kuis tanpa bantuan", "Mengukur pemisahan kinerja dan kompetensi", "Kuis singkat berkala dengan aturan jelas"],
            ["Jenis dukungan AI", "Membedakan jawaban langsung, petunjuk, umpan balik, dan penjelasan", "Kategori layanan dan jumlah langkah"],
            ["Interaksi teman", "Mengkalibrasi jaringan kontak dan bantuan", "Survei singkat siapa berdiskusi dengan siapa atau observasi agregat"],
            ["Usaha dan waktu", "Mengkalibrasi biaya yang memengaruhi adopsi", "Waktu tugas dan skala usaha singkat"],
            ["Pengetahuan awal dan retensi", "Mengkalibrasi heterogenitas dan pembelajaran", "Pretest dan delayed posttest"],
            ["Persepsi diri tervalidasi", "Mengganti proxy C dengan konstruk psikologis yang sah", "Skala self efficacy yang sesuai usia dan konteks"],
        ],
        [1.65, 2.55, 2.6],
        font_size=8.6,
    )
    add_body(
        doc,
        "Pengumpulan data harus melalui persetujuan institusi, persetujuan orang tua atau siswa sesuai usia, minimisasi data, pembatasan akses, dan pelaporan kelompok. Data jaringan sangat sensitif karena dapat mengungkap hubungan sosial meskipun nama dihapus.",
    )

    add_heading(doc, "Penilaian defensibilitas akhir", 1)
    add_body(
        doc,
        "Model defensibel sebagai eksperimen mekanistik transparan. Ia memiliki batas sebab akibat yang jelas, tidak memberi peran sosial kepada siswa, tidak memakai popularitas AI sebagai bonus, memisahkan jaringan kontak dari aliran jawaban, memasangkan dunia acak, mengunci prediksi, dan menguji sensitivitas. Keunggulan ini membuatnya lebih cocok sebagai baseline ilmiah daripada sistem agen bahasa yang sulit diaudit.",
    )
    add_body(
        doc,
        "Model belum defensibel sebagai alat ramalan atau evaluasi siswa. Parameter pendidikan belum dikalibrasi, jaringan bukan hasil pengukuran, tutor disederhanakan, dan angka adopsi sangat sensitif terhadap aturan keputusan. Klaim yang pantas adalah arah mekanisme bersyarat, bukan estimasi prevalensi atau efektivitas sekolah.",
    )
    add_body(
        doc,
        "Keputusan penelitian berikutnya sebaiknya bukan menambah kompleksitas sebanyak mungkin. Prioritasnya adalah menjaga baseline, menguji komponen agentik satu per satu, dan memastikan setiap peningkatan realisme menghasilkan pengukuran atau prediksi baru yang dapat dibantah.",
    )

    add_heading(doc, "Glosarium", 1)
    add_table(
        doc,
        ["Istilah", "Arti dalam laporan"],
        [
            ["Agen", "Satu siswa sintetis dengan keadaan dan aturan keputusan sendiri"],
            ["Simpul", "Representasi seorang siswa pada grafik jaringan"],
            ["Tepi atau garis", "Peluang kontak sosial antara dua siswa"],
            ["Emergence", "Pola tingkat kelas yang muncul dari interaksi lokal dan tidak ditetapkan sebagai tipe siswa"],
            ["Monte Carlo", "Mengulang simulasi dengan kejadian acak untuk melihat distribusi hasil"],
            ["Kalibrasi", "Menyesuaikan parameter agar cocok dengan data nyata yang relevan"],
            ["Validasi", "Memeriksa apakah model cukup sesuai untuk tujuan yang dinyatakan"],
            ["Counterfactual", "Dunia pembanding hipotetis yang berbeda pada kebijakan tertentu"],
            ["Scaffolding", "Bantuan bertahap yang mendukung siswa tanpa langsung mengambil alih seluruh tugas"],
            ["Homofili", "Kecenderungan individu yang mirip untuk terhubung"],
            ["Assortativitas", "Ukuran apakah simpul yang mirip cenderung saling terhubung"],
            ["Gini sumber", "Ukuran konsentrasi pasokan jawaban; makin tinggi berarti makin terkonsentrasi"],
        ],
        [1.65, 5.15],
    )

    add_heading(doc, "Referensi sejawat", 1)
    references = [
        ("1", "Bastani H, Bastani O, Sungu A, Ge H, Kabakci O, dan Mariman R. Generative AI without guardrails can harm learning. Proceedings of the National Academy of Sciences. 2025.", "https://doi.org/10.1073/pnas.2422633122"),
        ("2", "Kestin G, Miller K, Klales A, Milbourne T, dan kolega. AI tutoring outperforms in class active learning. Scientific Reports. 2025.", "https://doi.org/10.1038/s41598-025-97652-6"),
        ("3", "VanLehn K. The relative effectiveness of human tutoring intelligent tutoring systems and other tutoring systems. Educational Psychologist. 2011.", "https://doi.org/10.1080/00461520.2011.611369"),
        ("4", "Letourneau A dan kolega. A systematic review of AI driven intelligent tutoring systems in K to 12 education. npj Science of Learning. 2025.", "https://doi.org/10.1038/s41539-025-00320-7"),
        ("5", "Freeman S, Eddy SL, McDonough M, Smith MK, Okoroafor N, Jordt H, dan Wenderoth MP. Active learning increases student performance in science engineering and mathematics. PNAS. 2014.", "https://doi.org/10.1073/pnas.1319030111"),
        ("6", "Sweller J dan Cooper GA. The use of worked examples as a substitute for problem solving in learning algebra. Cognition and Instruction. 1985.", "https://doi.org/10.1207/s1532690xci0201_3"),
        ("7", "Roediger HL dan Karpicke JD. Test enhanced learning. Psychological Science. 2006.", "https://doi.org/10.1111/j.1467-9280.2006.01693.x"),
        ("8", "Hattie J dan Timperley H. The power of feedback. Review of Educational Research. 2007.", "https://doi.org/10.3102/003465430298487"),
        ("9", "Rendell L dan kolega. Why copy others. Insights from the social learning strategies tournament. Science. 2010.", "https://doi.org/10.1126/science.1184719"),
        ("10", "Nowak MA dan Sigmund K. Evolution of indirect reciprocity. Nature. 2005.", "https://doi.org/10.1038/nature04131"),
        ("11", "Watkins CJCH dan Dayan P. Q learning. Machine Learning. 1992.", "https://doi.org/10.1007/BF00992698"),
        ("12", "Holland PW, Laskey KB, dan Leinhardt S. Stochastic blockmodels first steps. Social Networks. 1983.", "https://doi.org/10.1016/0378-8733(83)90021-7"),
        ("13", "Grimm V dan kolega. The ODD protocol for describing agent based and other simulation models. Journal of Artificial Societies and Social Simulation. 2020.", "https://doi.org/10.18564/jasss.4259"),
        ("14", "Epstein JM. Why model. Journal of Artificial Societies and Social Simulation. 2008.", "https://www.jasss.org/11/4/12.html"),
        ("15", "Shalizi CR dan Thomas AC. Homophily and contagion are generically confounded in observational social network studies. Sociological Methods and Research. 2011.", "https://doi.org/10.1177/0049124111404820"),
        ("16", "McPherson M, Smith Lovin L, dan Cook JM. Birds of a feather. Homophily in social networks. Annual Review of Sociology. 2001.", "https://doi.org/10.1146/annurev.soc.27.1.415"),
        ("17", "Bruun J dan Brewe E. Talking and learning physics. Physical Review Special Topics Physics Education Research. 2013.", "https://doi.org/10.1103/PhysRevSTPER.9.020109"),
    ]
    for num, citation, url in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.23)
        p.paragraph_format.first_line_indent = Inches(-0.23)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"[{num}] {citation} ")
        set_run_font(r, size=9.4)
        add_hyperlink(p, "DOI atau artikel", url)

    add_heading(doc, "Catatan reproducibility", 1)
    add_body(
        doc,
        "Laporan ini dibuat dari keluaran beku E8. Angka utama berasal dari primary summaries dan primary trajectories. Grafik jaringan berasal dari log mentah replikasi 282 yang dipilih sebagai replikasi paling dekat dengan rata-rata multivariat. Tidak ada hasil simulasi baru yang dibuat untuk laporan ini.",
    )
    add_table(
        doc,
        ["Komponen", "Nilai"],
        [
            ["Versi model", "Classroom Swarm V3 dengan intervensi E8"],
            ["Jumlah run", "23.000"],
            ["Perbandingan utama", "400 replikasi per kebijakan"],
            ["Replikasi visual jaringan", "282 dipilih secara objektif"],
            ["Status kalibrasi", "Tidak ada kalibrasi data sekolah"],
            ["Status AI", "Layanan eksternal dan bukan simpul jaringan"],
        ],
        [2.2, 4.5],
    )

    # Core properties are intentionally generic; final QA also scrubs metadata.
    doc.core_properties.title = "Laporan metodologi simulasi jaringan kelas dengan akses AI"
    doc.core_properties.subject = "Metodologi dan hasil Classroom Swarm V3 E8"
    doc.core_properties.author = "Classroom Swarm Research Project"
    doc.core_properties.keywords = "agent based model classroom network AI education"
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    set_mpl()
    data = load_representative()
    means = primary_means()
    figures = {
        "anatomy": save_network_anatomy(data),
        "flows": save_flow_comparison(data),
        "primary": save_primary_results(means),
        "trajectory": save_trajectories(),
    }
    build_doc(figures, means)
    print(OUT)


if __name__ == "__main__":
    main()
